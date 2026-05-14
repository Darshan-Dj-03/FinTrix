from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(r"Z:\WORK\Extra-Projects\INTERNSHIP\Hostel-Management\FinTrix\Training_Report_MERN_Stack_Web_Development.docx")
DOMAIN_TITLE = "MERN Stack Web Development"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_run_font(run, size=12, bold=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_format(paragraph, before=0, after=8, line_spacing=1.5):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line_spacing


def add_styled_paragraph(document, text="", style=None, align=None, before=0, after=8):
    p = document.add_paragraph(style=style)
    if text:
        run = p.add_run(text)
        size = 12
        bold = False
        color = None
        if style == "Title":
            size = 20
            bold = True
            color = "1F4E78"
        elif style == "Heading 1":
            size = 16
            bold = True
            color = "1F1F1F"
        elif style == "Heading 2":
            size = 14
            bold = True
            color = "2F5597"
        set_run_font(run, size=size, bold=bold, color=color)
    if align is not None:
        p.alignment = align
    set_paragraph_format(p, before=before, after=after)
    return p


def add_body_paragraph(document, text, first_line_indent=0.3):
    p = document.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=12)
    set_paragraph_format(p, before=0, after=8)
    p.paragraph_format.first_line_indent = Inches(first_line_indent)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def add_bullet(document, text):
    p = document.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    set_run_font(run, size=12)
    set_paragraph_format(p, before=0, after=4)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def add_caption(document, text, kind="figure"):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=10, bold=True)
    set_paragraph_format(p, before=0, after=8, line_spacing=1.0)
    if kind == "table":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p


def set_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    set_run_font(run, size=10)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)


def apply_base_styles(document):
    for style_name, font_size in [("Normal", 12), ("Title", 20), ("Heading 1", 16), ("Heading 2", 14)]:
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(font_size)
    document.styles["Normal"].paragraph_format.line_spacing = 1.5


def configure_sections(document):
    for section in document.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(0.9)
        header = section.header
        header.is_linked_to_previous = False
        if header.paragraphs:
            header_p = header.paragraphs[0]
            header_p.clear()
        else:
            header_p = header.add_paragraph()
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header_run = header_p.add_run(DOMAIN_TITLE)
        set_run_font(header_run, size=10, bold=True, color="1F4E78")

        footer = section.footer
        footer.is_linked_to_previous = False
        footer_table = footer.add_table(rows=1, cols=2, width=Inches(6.3))
        footer_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        footer_table.columns[0].width = Inches(3.4)
        footer_table.columns[1].width = Inches(2.9)
        left_p = footer_table.cell(0, 0).paragraphs[0]
        left_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        left_run = left_p.add_run("Department of BCA")
        set_run_font(left_run, size=10)
        right_p = footer_table.cell(0, 1).paragraphs[0]
        set_page_number(right_p)
        for cell in footer_table.row_cells(0):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_cover_page(document):
    add_styled_paragraph(document, "Training Report", style="Title", align=WD_ALIGN_PARAGRAPH.CENTER, before=18, after=10)
    add_styled_paragraph(document, "On", style="Normal", align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=8)
    cover_subtitle = document.add_paragraph()
    cover_subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cover_subtitle.add_run(DOMAIN_TITLE)
    set_run_font(run, size=18, bold=True, color="2F5597")
    set_paragraph_format(cover_subtitle, before=0, after=18)

    info_table = document.add_table(rows=7, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.style = "Table Grid"
    labels = [
        ("Student Name", "[Enter Student Name]"),
        ("USN / Register Number", "[Enter USN]"),
        ("Course / Semester", "[Enter Course and Semester]"),
        ("Training Period", "[Enter Duration]"),
        ("Organization", "[Enter Organization Name]"),
        ("Academic Year", "[Enter Academic Year]"),
        ("Guide / Mentor", "[Enter Guide Name]"),
    ]
    for i, (label, value) in enumerate(labels):
        left = info_table.cell(i, 0)
        right = info_table.cell(i, 1)
        left.text = ""
        right.text = ""
        lp = left.paragraphs[0]
        rp = right.paragraphs[0]
        lr = lp.add_run(label)
        rr = rp.add_run(value)
        set_run_font(lr, size=12, bold=True)
        set_run_font(rr, size=12)
        left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(left, "D9EAF7")
    add_styled_paragraph(document, "", style="Normal", after=10)
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note_run = note.add_run(
        "Note: This document is prepared as a Training Report. It is separate from the Project Report "
        "and therefore focuses on learning activities, exercises, assignments, and technology exposure "
        "during the training period."
    )
    set_run_font(note_run, size=11, color="555555")
    set_paragraph_format(note, before=6, after=0)
    document.add_page_break()


def add_toc(document):
    add_styled_paragraph(document, "Table of Contents", style="Heading 1", before=0, after=10)
    toc_rows = [
        ("1", "Introduction", "1"),
        ("2", "Company Profile", "2"),
        ("3", "Objectives of the Internship Training", "3"),
        ("4", "Overview of MERN Stack", "4"),
        ("5", "Training Tasks, Exercises, and Assignments", "5"),
        ("6", "Detailed Learning Modules", "8"),
        ("7", "Week-Wise Training Log", "11"),
        ("8", "Skills Acquired", "13"),
        ("9", "Industry Practices and Professional Exposure", "15"),
        ("10", "Challenges Faced", "17"),
        ("11", "Key Learnings and Self Assessment", "18"),
        ("12", "Conclusion", "19"),
        ("13", "Appendix", "20"),
    ]
    table = document.add_table(rows=1 + len(toc_rows), cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Chapter No.", "Contents", "Page No."]
    for idx, title in enumerate(headers):
        cell = table.cell(0, idx)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title)
        set_run_font(r, size=11, bold=True)
        set_cell_shading(cell, "BFD7EE")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row_idx, row in enumerate(toc_rows, start=1):
        for col_idx, value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx != 1 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_run_font(r, size=11)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    add_caption(document, "Table 1: Chapter-Wise Contents of the Training Report", kind="table")
    document.add_page_break()


def add_company_profile_placeholder(document):
    add_styled_paragraph(document, "2. Company Profile", style="Heading 1", before=0, after=8)
    add_body_paragraph(
        document,
        "This section is intentionally left blank as requested and may be completed later with the "
        "company or organization details after approval from the guide or institution.",
        first_line_indent=0.0,
    )
    table = document.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    placeholders = [
        ("Company / Organization Name", "........................................................"),
        ("Location", "........................................................"),
        ("Year of Establishment", "........................................................"),
        ("Domain Expertise", "........................................................"),
        ("Assigned Department / Team", "........................................................"),
    ]
    for row_idx, (label, value) in enumerate(placeholders):
        left = table.cell(row_idx, 0)
        right = table.cell(row_idx, 1)
        left.text = ""
        right.text = ""
        lp = left.paragraphs[0]
        rp = right.paragraphs[0]
        lr = lp.add_run(label)
        rr = rp.add_run(value)
        set_run_font(lr, size=11, bold=True)
        set_run_font(rr, size=11)
        set_cell_shading(left, "EAF2F8")
    add_caption(document, "Table 2: Company Profile Details to Be Filled Later", kind="table")


def add_content(document):
    add_styled_paragraph(document, "1. Introduction", style="Heading 1", before=0, after=8)
    intro_paragraphs = [
        "Industrial training plays an important role in bridging the gap between academic learning and professional practice. During the training period, students are exposed to real development tools, coding standards, workflow discipline, and problem-solving methods that are difficult to understand fully through classroom theory alone. The present training report is prepared to document the learning journey, practical exercises, and technical understanding gained during the internship period in the domain of MERN Stack Web Development.",
        "The training was centered on the core technologies of modern JavaScript-based full stack development, namely MongoDB, Express.js, React.js, and Node.js. The overall purpose of the training was to develop a clear understanding of how frontend and backend technologies interact, how databases are integrated into web applications, and how developers use industry tools to plan, build, test, and improve software modules. Along with technical concepts, the training also helped in understanding teamwork, debugging workflow, documentation, code organization, and deployment basics.",
        "This report is different from a project report. It does not focus on a single final product or project outcome. Instead, it focuses on the training activities carried out throughout the internship period, such as practice assignments, interface-building exercises, API development tasks, database operations, version control activities, testing efforts, and module-wise implementation practice. The report therefore captures the knowledge development process and the gradual strengthening of practical skills during the training period.",
        "The expectations from the training were to gain hands-on familiarity with industry-relevant tools, understand real-world software development practices, improve confidence in working with a complete technology stack, and learn how to translate theoretical knowledge into structured coding tasks. The training also aimed to build readiness for future professional work by encouraging consistency, logical thinking, collaboration, and the habit of writing maintainable solutions."
    ]
    for para in intro_paragraphs:
        add_body_paragraph(document, para)

    add_company_profile_placeholder(document)

    add_styled_paragraph(document, "3. Objectives of the Internship Training", style="Heading 1", before=12, after=8)
    objective_paragraph = (
        "The internship training was undertaken with the objective of obtaining practical exposure to the "
        "software development lifecycle and strengthening understanding of full stack web application "
        "development using the MERN ecosystem. The major objectives identified during the training period were:"
    )
    add_body_paragraph(document, objective_paragraph, first_line_indent=0.0)
    objectives = [
        "To understand how a real development environment is organized with source control, coding tasks, testing, and version updates.",
        "To learn the fundamental concepts and workflow of MongoDB, Express.js, React.js, and Node.js through guided exercises and daily assignments.",
        "To practice building user interfaces, reusable components, forms, and state-driven screens using React.js.",
        "To understand backend routing, middleware, request handling, authentication basics, and API response patterns using Express.js and Node.js.",
        "To gain familiarity with database schema design, CRUD operations, validation, and storage of structured data through MongoDB.",
        "To improve debugging skills, error tracing, and problem-solving techniques while integrating frontend and backend modules.",
        "To develop discipline in documentation, code readability, and collaborative practices such as Git-based version control and review."
    ]
    for item in objectives:
        add_bullet(document, item)

    add_styled_paragraph(document, "4. Overview of MERN Stack", style="Heading 1", before=12, after=8)
    overview_intro = (
        "MERN Stack is a popular full stack JavaScript technology combination used to build responsive, "
        "database-driven, and scalable web applications. It allows developers to use JavaScript across both "
        "client-side and server-side development, which makes implementation more consistent and easier to manage."
    )
    add_body_paragraph(document, overview_intro, first_line_indent=0.0)

    tech_items = [
        (
            "MongoDB",
            "MongoDB is a NoSQL document-oriented database used to store application data in flexible JSON-like documents. During the training period, it was studied for schema design, collections, CRUD operations, and connection handling through backend code."
        ),
        (
            "Express.js",
            "Express.js is a backend framework built on Node.js. It was used to understand routing, middleware, request validation, error handling, and the creation of REST-style APIs for communication between the client and server."
        ),
        (
            "React.js",
            "React.js is a frontend library used for building interactive user interfaces. The training included component creation, props, state management, form handling, event-driven UI updates, and modular screen development."
        ),
        (
            "Node.js",
            "Node.js is a JavaScript runtime that enables server-side development. It was used to run backend applications, manage packages, connect APIs, and execute application logic in a scalable environment."
        ),
    ]
    for heading, body in tech_items:
        add_styled_paragraph(document, heading, style="Heading 2", before=6, after=4)
        add_body_paragraph(document, body, first_line_indent=0.0)

    add_styled_paragraph(document, "5. Training Tasks, Exercises, and Assignments", style="Heading 1", before=12, after=8)
    task_intro = (
        "The training report emphasizes practice-oriented activities carried out during the internship period. "
        "The following tasks were not treated as final project submissions; rather, they were learning exercises "
        "designed to improve command over different parts of the MERN Stack."
    )
    add_body_paragraph(document, task_intro, first_line_indent=0.0)

    training_tasks = [
        (
            "Basic Environment Setup and Tool Familiarization",
            "The initial phase of training involved installation and configuration of Node.js, package managers, code editors, Git, and browser developer tools. Exercises included creating sample folders, running starter applications, understanding dependency files, and using terminal commands for routine development activities."
        ),
        (
            "React Component Practice",
            "Several small UI assignments were performed to understand JSX structure, reusable components, props, conditional rendering, and form interaction. Simple screens such as login forms, list displays, table layouts, modal views, and status cards were created and refined during practice sessions."
        ),
        (
            "State Management and Form Handling",
            "Assignments were carried out to manage user input, validate form data, update views dynamically, and reset or submit state-driven forms. These exercises helped in understanding how React handles live data changes and improves user interaction in web interfaces."
        ),
        (
            "Backend Routing and API Creation",
            "Training tasks included creating Express routes, building API endpoints for create, read, update, and delete operations, structuring controllers, and testing request-response flow. API testing was practiced using Postman to verify payload handling, error messages, and status codes."
        ),
        (
            "Database Connectivity and CRUD Operations",
            "Practice modules were completed to connect the backend with MongoDB, define schemas, store data records, retrieve results, update fields, and delete entries safely. This part of the training improved understanding of database relationships, validation, and model organization."
        ),
        (
            "Authentication and Security Basics",
            "Simple assignments were used to understand login flow, protected routes, token-based access concepts, password handling basics, and validation of user identity before allowing access to selected features."
        ),
        (
            "Debugging and Error Resolution Exercises",
            "During training, common problems such as failed API calls, incorrect routes, state update issues, undefined variables, and database connection errors were intentionally analyzed and resolved. This improved confidence in reading logs, tracing bugs, and correcting implementation mistakes."
        ),
        (
            "Version Control and Collaboration Practice",
            "Git and GitHub were used to understand repository structure, commits, branch-based work, code updates, and tracking of changes over time. Training assignments required maintaining discipline in saving progress and documenting updates in a structured way."
        ),
        (
            "Deployment Awareness and Build Verification",
            "Basic deployment-oriented tasks such as preparing environment configuration, checking production builds, understanding hosting flow, and reviewing deployment-related issues were included to provide exposure beyond local development."
        ),
    ]
    for title, body in training_tasks:
        add_styled_paragraph(document, title, style="Heading 2", before=6, after=4)
        add_body_paragraph(document, body, first_line_indent=0.0)

    task_table = document.add_table(rows=1 + 6, cols=4)
    task_table.style = "Table Grid"
    task_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    task_headers = ["Sl. No.", "Training Activity", "Tools / Technologies", "Learning Outcome"]
    for idx, value in enumerate(task_headers):
        cell = task_table.cell(0, idx)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value)
        set_run_font(r, size=10, bold=True)
        set_cell_shading(cell, "D9EAF7")
    task_rows = [
        ("1", "Setup and package management practice", "Node.js, npm, VS Code", "Understood environment preparation and dependency handling"),
        ("2", "UI and form-building exercises", "React.js, JSX, CSS", "Improved component design and form interaction skills"),
        ("3", "REST API practice", "Express.js, Postman", "Learned routing, requests, and structured responses"),
        ("4", "Database CRUD assignments", "MongoDB, Mongoose", "Understood schema design and data operations"),
        ("5", "Authentication basics", "Node.js, JWT concepts", "Learned protected access and validation flow"),
        ("6", "Version control practice", "Git, GitHub", "Developed discipline in tracking and collaborating on code"),
    ]
    for row_idx, row in enumerate(task_rows, start=1):
        for col_idx, value in enumerate(row):
            cell = task_table.cell(row_idx, col_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx != 1 and col_idx != 3 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_run_font(r, size=10)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    add_caption(document, "Table 3: Summary of Training Activities and Outcomes", kind="table")

    add_styled_paragraph(document, "6. Detailed Learning Modules", style="Heading 1", before=12, after=8)
    detailed_intro = (
        "Apart from broad tasks, the training period was divided into focused learning modules so that each "
        "technology area could be understood progressively. These modules were useful because they broke down "
        "complex full stack development into smaller and manageable stages of learning."
    )
    add_body_paragraph(document, detailed_intro, first_line_indent=0.0)

    detailed_modules = [
        (
            "6.1 Frontend Layout and Styling Exercises",
            [
                "Training began with layout-oriented exercises such as creating responsive sections, navigation bars, cards, tables, form containers, and dashboard-style views. These activities helped in understanding HTML structure through JSX, CSS-based alignment, spacing consistency, and reusable UI design patterns.",
                "Special attention was given to how a user interface should remain readable and structured across different sections of a page. Simple design corrections such as label alignment, button hierarchy, spacing between fields, and clarity of headings were practiced repeatedly to develop a better visual sense for application interfaces.",
            ],
        ),
        (
            "6.2 React State, Props, and Component Communication",
            [
                "After gaining familiarity with layout creation, the training focused on making interfaces interactive. Exercises were carried out to pass data through props, update state after user actions, conditionally render sections, and manage dynamic values inside forms and lists.",
                "This module was especially helpful in understanding how a UI changes in response to input and why component separation improves maintainability. Through repeated practice, it became easier to divide a large screen into smaller reusable parts and connect them with predictable data flow.",
            ],
        ),
        (
            "6.3 Form Validation and User Input Handling",
            [
                "A separate set of assignments was used to practice text fields, dropdowns, date inputs, validation messages, and submission workflows. The goal was to understand how to prevent invalid input, how to show feedback to users, and how to reset or preserve values when needed.",
                "This module improved awareness of user experience because form handling is not only about accepting data, but also about guiding the user toward correct and complete submission. It also helped in understanding why careful validation is essential before backend processing begins.",
            ],
        ),
        (
            "6.4 API Request and Response Management",
            [
                "Frontend-to-backend communication was studied through modules that involved sending requests, receiving responses, showing loading indicators, and handling failures. Assignments were framed around practical actions such as fetching lists, creating sample records, editing stored data, and confirming updates from the server.",
                "These exercises improved understanding of asynchronous behavior and helped in learning where errors can appear during communication. They also built confidence in connecting screens to real API endpoints instead of static mock data.",
            ],
        ),
        (
            "6.5 Backend Logic and Controller Organization",
            [
                "The backend portion of training included writing route handlers, separating logic into controller functions, validating requests, and organizing responses in a consistent structure. Sample exercises included creating modules for registration-style flows, list retrieval, updates, and record deletion.",
                "This module clarified the importance of clean backend structure. It showed that readable controller code, reusable helper methods, and proper validation improve maintenance and reduce confusion when multiple endpoints are involved.",
            ],
        ),
        (
            "6.6 Database Schema and Record Management",
            [
                "Schema design exercises helped in understanding how data fields are structured, how required values are enforced, and how related information can be stored safely. CRUD practice was carried out repeatedly so that the process of adding, reading, modifying, and deleting records could be understood thoroughly.",
                "The module also highlighted the importance of field naming consistency and validation logic at the data layer. These learnings are essential because weak schema planning can create errors throughout the rest of the application.",
            ],
        ),
        (
            "6.7 Authentication and Protected Workflow Understanding",
            [
                "During the authentication module, simple protected workflow patterns were studied. Exercises focused on understanding login concepts, token-based identity verification, restricted actions, and the difference between public routes and authenticated routes.",
                "Although this was handled at a training level rather than through a major product implementation, it helped in understanding the need for access control, request authorization, and security-minded development habits.",
            ],
        ),
        (
            "6.8 Build, Testing, and Deployment Readiness",
            [
                "The later phase of training introduced production awareness. This included checking whether frontend code builds successfully, understanding environment variables, reviewing server startup behavior, and observing what must be configured properly before deployment.",
                "Even basic exposure to these checks made an important difference because it showed that software development is not complete when code works locally. It must also be stable, testable, and reproducible in other environments.",
            ],
        ),
    ]
    for heading, paragraphs in detailed_modules:
        add_styled_paragraph(document, heading, style="Heading 2", before=6, after=4)
        for paragraph in paragraphs:
            add_body_paragraph(document, paragraph, first_line_indent=0.0)

    module_table = document.add_table(rows=1 + 8, cols=4)
    module_table.style = "Table Grid"
    module_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    module_headers = ["Module", "Main Focus", "Practice Type", "Outcome"]
    for idx, value in enumerate(module_headers):
        cell = module_table.cell(0, idx)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value)
        set_run_font(r, size=10, bold=True)
        set_cell_shading(cell, "D9EAF7")
    module_rows = [
        ("1", "UI layout and styling", "Screens, forms, cards", "Improved presentation and structure"),
        ("2", "State and props", "Interactive component exercises", "Better control of UI updates"),
        ("3", "Validation", "Input checking and error display", "More reliable form handling"),
        ("4", "API integration", "Request-response exercises", "Clear understanding of async flow"),
        ("5", "Backend logic", "Routes and controllers", "Structured server-side implementation"),
        ("6", "Database work", "Schema and CRUD practice", "Better understanding of persistence"),
        ("7", "Authentication", "Protected-flow study", "Awareness of access control"),
        ("8", "Deployment checks", "Build and environment review", "Industry readiness mindset"),
    ]
    for row_idx, row in enumerate(module_rows, start=1):
        for col_idx, value in enumerate(row):
            cell = module_table.cell(row_idx, col_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_run_font(r, size=10)
    add_caption(document, "Table 4: Module-Wise Learning Focus During Training", kind="table")

    add_styled_paragraph(document, "7. Week-Wise Training Log", style="Heading 1", before=12, after=8)
    week_intro = (
        "A week-wise view of the training process helps show how the learning journey progressed from basic "
        "tool setup to more structured full stack practice. The following log summarizes the flow of activities "
        "carried out during the training period."
    )
    add_body_paragraph(document, week_intro, first_line_indent=0.0)

    week_table = document.add_table(rows=1 + 8, cols=4)
    week_table.style = "Table Grid"
    week_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    week_headers = ["Week", "Topics Covered", "Practice / Assignment", "Learning Summary"]
    for idx, value in enumerate(week_headers):
        cell = week_table.cell(0, idx)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value)
        set_run_font(r, size=10, bold=True)
        set_cell_shading(cell, "CFE2F3")
    week_rows = [
        ("1", "Environment setup and tool orientation", "Installed tools, ran starter apps", "Became familiar with local development workflow"),
        ("2", "JSX, components, and UI basics", "Built sample screens and cards", "Understood reusable frontend structure"),
        ("3", "Forms and state handling", "Created input-driven forms", "Improved event handling and validation"),
        ("4", "API concepts and Postman testing", "Created and tested sample endpoints", "Learned how data moves between client and server"),
        ("5", "MongoDB and schema design", "Performed CRUD practice", "Understood persistent storage and validation"),
        ("6", "Authentication concepts", "Observed token-based access flow", "Gained awareness of protected operations"),
        ("7", "Code organization and debugging", "Resolved structured issues", "Improved tracing and correction methods"),
        ("8", "Build and deployment awareness", "Reviewed configuration and builds", "Understood production-readiness checks"),
    ]
    for row_idx, row in enumerate(week_rows, start=1):
        for col_idx, value in enumerate(row):
            cell = week_table.cell(row_idx, col_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_run_font(r, size=10)
    add_caption(document, "Table 5: Week-Wise Training Progress Log", kind="table")

    weekly_reflections = [
        "The week-wise approach made the training more systematic because each stage built naturally on the previous one. Early familiarity with tools reduced confusion in later assignments and created a stable base for practical work.",
        "The progression from interface-level practice to backend handling and then to integration exercises also made the learning process more meaningful. Instead of studying technologies in isolation, the training gradually demonstrated how all layers must work together to produce a functioning web application.",
        "By the final phase, the training experience had moved beyond isolated tasks and had developed into a broader understanding of workflow, debugging discipline, documentation, and deployment awareness, which are all important for real software development environments.",
    ]
    for paragraph in weekly_reflections:
        add_body_paragraph(document, paragraph, first_line_indent=0.0)

    add_styled_paragraph(document, "8. Skills Acquired", style="Heading 1", before=12, after=8)
    add_styled_paragraph(document, "Technical Skills", style="Heading 2", before=4, after=4)
    technical_skills = [
        "Understanding of full stack development flow using MongoDB, Express.js, React.js, and Node.js.",
        "Hands-on exposure to API design, client-server communication, and CRUD-based data handling.",
        "Working knowledge of frontend forms, reusable components, event handling, and UI state updates.",
        "Basic understanding of authentication, validation, configuration management, and deployment checks.",
    ]
    for item in technical_skills:
        add_bullet(document, item)
    technical_paragraphs = [
        "The most valuable technical gain from the training was the ability to understand the complete path of data flow in a full stack application. This includes the journey from user interaction in the frontend, to API communication, to backend handling, and finally to database storage or retrieval.",
        "Another major improvement was the ability to read and organize application code more effectively. Through repeated exposure to components, routes, controllers, and models, the training strengthened the habit of separating concerns and maintaining cleaner project structure.",
    ]
    for paragraph in technical_paragraphs:
        add_body_paragraph(document, paragraph, first_line_indent=0.0)
    add_styled_paragraph(document, "Tools and Software", style="Heading 2", before=4, after=4)
    tools_skills = [
        "VS Code for code editing and project organization.",
        "Git and GitHub for source control and update tracking.",
        "Postman for API testing and request verification.",
        "MongoDB Atlas or local MongoDB tools for data storage and inspection.",
    ]
    for item in tools_skills:
        add_bullet(document, item)
    add_body_paragraph(
        document,
        "Regular use of these tools also improved comfort with a practical developer environment. Instead of treating tools as separate utilities, the training demonstrated how editors, version control, API clients, and databases form a connected daily workflow.",
        first_line_indent=0.0,
    )
    add_styled_paragraph(document, "Soft Skills", style="Heading 2", before=4, after=4)
    soft_skills = [
        "Improved communication while discussing issues and understanding assigned tasks.",
        "Better time management through task-wise completion and progress tracking.",
        "Increased patience and analytical thinking while debugging and resolving errors.",
        "Greater confidence in learning new frameworks and working in a structured development environment.",
    ]
    for item in soft_skills:
        add_bullet(document, item)
    add_body_paragraph(
        document,
        "Soft-skill development was not separate from technical growth; it happened alongside it. Repeated debugging, revision of assignments, and structured completion of tasks created a more disciplined and patient approach toward problem solving.",
        first_line_indent=0.0,
    )

    add_styled_paragraph(document, "9. Industry Practices and Professional Exposure", style="Heading 1", before=12, after=8)
    industry_paragraphs = [
        "One of the strongest outcomes of the training period was exposure to the discipline expected in a professional development setting. Even when the work involved small practice modules, the training emphasized structured folders, naming consistency, readable code, and step-wise implementation rather than unplanned coding.",
        "Version control practices introduced the importance of tracking changes responsibly. It became clear that software development is not only about writing code but also about documenting progress, maintaining traceable updates, and making collaboration easier for other team members.",
        "The training also highlighted the importance of testing and review. A feature or screen should not be considered complete just because it appears to work once; instead, it should be observed under different conditions, checked for invalid input, and verified for response quality. This mindset reflects an industry-oriented approach to quality.",
        "Another useful professional exposure was understanding that deployment and runtime behavior require separate attention. Environment values, build configuration, server start-up handling, and production readiness checks can affect application stability even after local development is complete. This awareness is essential for moving from student practice to production thinking.",
    ]
    for paragraph in industry_paragraphs:
        add_body_paragraph(document, paragraph, first_line_indent=0.0)

    practice_table = document.add_table(rows=1 + 5, cols=3)
    practice_table.style = "Table Grid"
    practice_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    practice_headers = ["Practice Area", "Observed Importance", "Training Value"]
    for idx, value in enumerate(practice_headers):
        cell = practice_table.cell(0, idx)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value)
        set_run_font(r, size=10, bold=True)
        set_cell_shading(cell, "D9EAD3")
    practice_rows = [
        ("Code organization", "Improves readability and maintenance", "Built cleaner implementation habits"),
        ("Version control", "Supports collaboration and recovery", "Improved update discipline"),
        ("Testing", "Reduces errors before release", "Strengthened verification mindset"),
        ("Documentation", "Clarifies work for future reference", "Improved reporting and clarity"),
        ("Deployment awareness", "Prepares code for real environments", "Connected learning with production context"),
    ]
    for row_idx, row in enumerate(practice_rows, start=1):
        for col_idx, value in enumerate(row):
            cell = practice_table.cell(row_idx, col_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_run_font(r, size=10)
    add_caption(document, "Table 6: Professional Practices Observed During Training", kind="table")

    add_styled_paragraph(document, "10. Challenges Faced", style="Heading 1", before=12, after=8)
    challenges = [
        "Understanding the integration between frontend components, backend routes, and database records was initially challenging because each layer depends on the others for correct operation.",
        "API testing required careful attention to request payloads, route paths, and error responses, especially when the structure of incoming data was not aligned with backend expectations.",
        "Managing state updates and form-driven UI behavior in React required repeated practice to avoid stale values, incorrect rendering, or incomplete validation.",
        "Debugging runtime and configuration issues, including package dependencies, server responses, and environment variables, required patience and regular use of logs and developer tools.",
        "These challenges were gradually resolved by practicing smaller assignments, reviewing errors step by step, testing each module independently, and taking guidance when necessary."
    ]
    for challenge in challenges:
        add_body_paragraph(document, challenge, first_line_indent=0.0)
    challenge_resolution = [
        "A useful lesson from these difficulties was that complex problems become easier when broken into smaller checks. Verifying one layer at a time, such as form values, request payload, route behavior, or database response, reduced confusion and saved effort during debugging.",
        "The training also reinforced the importance of calm and methodical error handling. Instead of making repeated random changes, it proved more effective to inspect logs, compare expected and actual data, and test corrections in a controlled order.",
    ]
    for paragraph in challenge_resolution:
        add_body_paragraph(document, paragraph, first_line_indent=0.0)

    add_styled_paragraph(document, "11. Key Learnings and Self Assessment", style="Heading 1", before=12, after=8)
    key_learning_intro = (
        "At the end of the training period, it was possible to identify several important learning outcomes that "
        "went beyond individual assignments. These outcomes reflect both technical growth and readiness for more "
        "structured project responsibilities in the future."
    )
    add_body_paragraph(document, key_learning_intro, first_line_indent=0.0)

    key_learning_points = [
        "Full stack development becomes easier to understand when it is studied as a connected flow rather than as separate technologies.",
        "Small and repeated exercises are highly effective for building confidence with UI logic, API handling, and database operations.",
        "Debugging is a skill that improves through observation, patience, and disciplined testing rather than by trial and error alone.",
        "Readable code, proper naming, and structured files make a strong difference in how quickly a module can be maintained or improved.",
        "Learning industry tools such as Git, Postman, and build workflows is essential because software development depends on process as much as coding ability.",
    ]
    for point in key_learning_points:
        add_bullet(document, point)

    self_assessment_paragraphs = [
        "From a self-assessment perspective, the training period significantly improved confidence in reading existing code and understanding how different layers of a web application communicate. Tasks that initially appeared complex became more manageable after repeated exposure to components, request flow, and data handling.",
        "The training also created awareness of areas that still require further improvement, such as advanced state management, deeper security implementation, testing automation, and deployment troubleshooting. Recognizing these improvement areas is useful because it gives direction for future learning beyond the training period.",
        "Overall, the training created a strong foundation and also clarified the next steps for skill development. This combination of present understanding and future learning direction is one of the most meaningful results of the internship experience.",
    ]
    for paragraph in self_assessment_paragraphs:
        add_body_paragraph(document, paragraph, first_line_indent=0.0)

    self_table = document.add_table(rows=1 + 6, cols=3)
    self_table.style = "Table Grid"
    self_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    self_headers = ["Area", "Training Level Achieved", "Future Improvement Focus"]
    for idx, value in enumerate(self_headers):
        cell = self_table.cell(0, idx)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value)
        set_run_font(r, size=10, bold=True)
        set_cell_shading(cell, "F4CCCC")
    self_rows = [
        ("React components", "Good working understanding", "Advanced reusable patterns"),
        ("Form handling", "Comfortable with basic validation", "Complex and dynamic validation"),
        ("API integration", "Able to test and connect endpoints", "Error resilience and optimization"),
        ("Backend logic", "Understood route-controller structure", "Advanced middleware design"),
        ("Database handling", "Comfortable with CRUD and schemas", "Optimization and data modeling"),
        ("Deployment awareness", "Basic build and config knowledge", "Production troubleshooting"),
    ]
    for row_idx, row in enumerate(self_rows, start=1):
        for col_idx, value in enumerate(row):
            cell = self_table.cell(row_idx, col_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_run_font(r, size=10)
    add_caption(document, "Table 7: Self Assessment of Learning Progress During Training", kind="table")

    add_styled_paragraph(document, "12. Conclusion", style="Heading 1", before=12, after=8)
    conclusion_paragraphs = [
        "The internship training provided valuable practical exposure to the MERN Stack domain and helped transform theoretical classroom knowledge into meaningful technical understanding. Through a sequence of guided tasks, exercises, and coding assignments, the training period improved confidence in building interfaces, handling APIs, connecting databases, and managing the logic of full stack applications.",
        "Apart from technical knowledge, the training also improved professional qualities such as consistency, patience, communication, and the ability to approach errors in a structured manner. The experience highlighted the importance of writing organized code, testing carefully, and understanding how multiple technologies work together in real development environments.",
        "Another important outcome was the realization that technology learning is most effective when theory, experimentation, and revision are combined. The training period demonstrated that steady practice across small but meaningful tasks leads to deeper understanding than passive observation alone.",
        "Overall, the training period served as an important foundation for future software development work. It strengthened interest in web technologies, prepared the learner for more advanced development tasks, and contributed positively toward long-term academic and career goals in the field of software engineering."
    ]
    for para in conclusion_paragraphs:
        add_body_paragraph(document, para)

    add_styled_paragraph(document, "13. Appendix", style="Heading 1", before=12, after=8)
    appendix_intro = (
        "Only training-related screenshots and supporting materials should be attached in this section. "
        "Project-specific screenshots are intentionally excluded from this Training Report."
    )
    add_body_paragraph(document, appendix_intro, first_line_indent=0.0)

    appendix_items = [
        ("Figure 1", "Screenshot of environment setup or folder structure created during training"),
        ("Figure 2", "Screenshot of React component or form exercise"),
        ("Figure 3", "Screenshot of API testing in Postman"),
        ("Figure 4", "Screenshot of MongoDB collection or CRUD exercise"),
        ("Figure 5", "Screenshot of Git commit history or version control activity"),
        ("Figure 6", "Screenshot of state management or validation practice"),
        ("Figure 7", "Screenshot of backend route or server response exercise"),
        ("Figure 8", "Screenshot of deployment or build verification activity"),
        ("Figure 9", "Screenshot of reusable component practice"),
        ("Figure 10", "Screenshot of sample table or dashboard exercise"),
        ("Figure 11", "Screenshot of authentication or protected route practice"),
        ("Figure 12", "Screenshot of debugging or console log analysis"),
        ("Figure 13", "Screenshot of package installation or terminal workflow"),
        ("Figure 14", "Screenshot of API payload verification"),
        ("Figure 15", "Screenshot of schema definition or model practice"),
        ("Figure 16", "Screenshot of training assignment notes or task checklist"),
        ("Figure 17", "Screenshot of component props or data passing exercise"),
        ("Figure 18", "Screenshot of conditional rendering practice"),
        ("Figure 19", "Screenshot of form validation message handling"),
        ("Figure 20", "Screenshot of edit or update workflow exercise"),
        ("Figure 21", "Screenshot of delete confirmation or list refresh practice"),
        ("Figure 22", "Screenshot of middleware or request validation study"),
        ("Figure 23", "Screenshot of local server start-up and request handling"),
        ("Figure 24", "Screenshot of code organization or folder modularization"),
        ("Figure 25", "Screenshot of error-tracing notes during debugging"),
        ("Figure 26", "Screenshot of production build or deployment-ready output"),
    ]
    for index, (fig_no, description) in enumerate(appendix_items, start=1):
        box = document.add_table(rows=1, cols=1)
        box.alignment = WD_TABLE_ALIGNMENT.CENTER
        box.style = "Table Grid"
        cell = box.cell(0, 0)
        cell.text = ""
        cell.width = Inches(5.8)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("\nInsert Training Screenshot Here\n")
        set_run_font(r, size=12, bold=True, color="7F8C8D")
        p2 = cell.add_paragraph(description)
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p2.runs[0], size=10)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        add_caption(document, f"{fig_no}: {description}", kind="figure")
        if index in {8, 16, 20}:
            document.add_page_break()

    add_styled_paragraph(document, "Appendix Notes", style="Heading 2", before=10, after=4)
    appendix_notes = [
        "All screenshots attached to the appendix should come only from training exercises, technical practice sessions, tool setup activities, or assignment-based learning modules.",
        "Screenshots from the final hostel-related project should be avoided in this report because the Training Report and Project Report are separate documents with different purposes.",
        "If required, each screenshot can be accompanied by a short explanation mentioning what was practiced, what technology was involved, and what was learned from that activity.",
    ]
    for note in appendix_notes:
        add_body_paragraph(document, note, first_line_indent=0.0)

    support_table = document.add_table(rows=1 + 8, cols=3)
    support_table.style = "Table Grid"
    support_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    support_headers = ["Sl. No.", "Supporting Material", "Purpose in Training Report"]
    for idx, value in enumerate(support_headers):
        cell = support_table.cell(0, idx)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value)
        set_run_font(r, size=10, bold=True)
        set_cell_shading(cell, "FFF2CC")
    support_rows = [
        ("1", "Tool setup screenshot", "Shows readiness of development environment"),
        ("2", "Component practice screenshot", "Demonstrates frontend exercise work"),
        ("3", "Form validation screenshot", "Shows input handling and feedback practice"),
        ("4", "Postman response screenshot", "Represents API verification activity"),
        ("5", "Database collection screenshot", "Supports CRUD and schema learning"),
        ("6", "Git activity screenshot", "Shows version control exposure"),
        ("7", "Build verification screenshot", "Confirms production-awareness exercise"),
        ("8", "Assignment notes or checklist", "Documents training progression"),
    ]
    for row_idx, row in enumerate(support_rows, start=1):
        for col_idx, value in enumerate(row):
            cell = support_table.cell(row_idx, col_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_run_font(r, size=10)
    add_caption(document, "Table 8: Suggested Supporting Materials for the Appendix", kind="table")


def main():
    document = Document()
    apply_base_styles(document)
    configure_sections(document)
    add_cover_page(document)
    add_toc(document)
    add_content(document)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()

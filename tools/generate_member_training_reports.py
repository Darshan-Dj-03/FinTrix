from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"Z:\WORK\Extra-Projects\INTERNSHIP\Hostel-Management\FinTrix")
DOMAIN_TITLE = "MERN Stack Web Development"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_run_font(run, size=12, bold=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_format(paragraph, before=0, after=8, line_spacing=1.5):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line_spacing


def add_styled_paragraph(document, text="", style=None, align=None, before=0, after=8):
    p = document.add_paragraph(style=style)
    if text:
        run = p.add_run(text)
        size = 12
        bold = False
        color = None
        if style == "Title":
            size = 20
            bold = True
            color = "1F4E78"
        elif style == "Heading 1":
            size = 16
            bold = True
            color = "1F1F1F"
        elif style == "Heading 2":
            size = 14
            bold = True
            color = "2F5597"
        set_run_font(run, size=size, bold=bold, color=color)
    if align is not None:
        p.alignment = align
    set_paragraph_format(p, before=before, after=after)
    return p


def add_body_paragraph(document, text, first_line_indent=0.3):
    p = document.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=12)
    set_paragraph_format(p, before=0, after=8)
    p.paragraph_format.first_line_indent = Inches(first_line_indent)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def add_bullet(document, text):
    p = document.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    set_run_font(run, size=12)
    set_paragraph_format(p, before=0, after=4)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def add_caption(document, text, kind="figure"):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=10, bold=True)
    set_paragraph_format(p, before=0, after=8, line_spacing=1.0)
    if kind == "table":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p


def set_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    set_run_font(run, size=10)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)


def apply_base_styles(document):
    for style_name, font_size in [("Normal", 12), ("Title", 20), ("Heading 1", 16), ("Heading 2", 14)]:
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(font_size)
    document.styles["Normal"].paragraph_format.line_spacing = 1.5


def configure_sections(document):
    for section in document.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(0.9)

        header = section.header
        header.is_linked_to_previous = False
        header_p = header.paragraphs[0]
        header_p.clear()
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header_run = header_p.add_run(DOMAIN_TITLE)
        set_run_font(header_run, size=10, bold=True, color="1F4E78")

        footer = section.footer
        footer.is_linked_to_previous = False
        footer_table = footer.add_table(rows=1, cols=2, width=Inches(6.3))
        footer_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        left_p = footer_table.cell(0, 0).paragraphs[0]
        left_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        left_run = left_p.add_run("Department of BCA")
        set_run_font(left_run, size=10)
        right_p = footer_table.cell(0, 1).paragraphs[0]
        set_page_number(right_p)
        for cell in footer_table.row_cells(0):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_cover_page(document, member_name, report_title):
    add_styled_paragraph(document, "Training Report", style="Title", align=WD_ALIGN_PARAGRAPH.CENTER, before=18, after=10)
    add_styled_paragraph(document, "On", style="Normal", align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=8)
    cover_subtitle = document.add_paragraph()
    cover_subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cover_subtitle.add_run(report_title)
    set_run_font(run, size=18, bold=True, color="2F5597")
    set_paragraph_format(cover_subtitle, before=0, after=18)

    info_table = document.add_table(rows=8, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.style = "Table Grid"
    labels = [
        ("Student Name", member_name),
        ("USN / Register Number", "[Enter USN]"),
        ("Course / Semester", "[Enter Course and Semester]"),
        ("Training Period", "[Enter Duration]"),
        ("Organization", "[Enter Organization Name]"),
        ("Academic Year", "[Enter Academic Year]"),
        ("Technology Domain", DOMAIN_TITLE),
        ("Guide / Mentor", "[Enter Guide Name]"),
    ]
    for i, (label, value) in enumerate(labels):
        left = info_table.cell(i, 0)
        right = info_table.cell(i, 1)
        left.text = ""
        right.text = ""
        lp = left.paragraphs[0]
        rp = right.paragraphs[0]
        lr = lp.add_run(label)
        rr = rp.add_run(value)
        set_run_font(lr, size=12, bold=True)
        set_run_font(rr, size=12)
        set_cell_shading(left, "D9EAF7")
        left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note_run = note.add_run(
        "Note: This document is prepared as an individual Training Report and focuses on learning activities, "
        "module-specific exercises, assignments, and workflow understanding gained during the training period."
    )
    set_run_font(note_run, size=11, color="555555")
    set_paragraph_format(note, before=10, after=0)
    document.add_page_break()


def add_toc(document):
    add_styled_paragraph(document, "Table of Contents", style="Heading 1", before=0, after=10)
    toc_rows = [
        ("1", "Introduction", "1"),
        ("2", "Company Profile", "2"),
        ("3", "Objectives of the Internship Training", "3"),
        ("4", "Technology and Domain Overview", "4"),
        ("5", "Member Module Allocation", "6"),
        ("6", "Detailed Module Study", "8"),
        ("7", "Module Integration and Workflow Mapping", "14"),
        ("8", "Week-Wise Training Log", "17"),
        ("9", "Skills Acquired", "20"),
        ("10", "Industry Practices and Professional Exposure", "22"),
        ("11", "Challenges Faced and Resolution Approach", "24"),
        ("12", "Key Learnings and Self Assessment", "25"),
        ("13", "Conclusion", "27"),
        ("14", "Appendix", "28"),
    ]
    table = document.add_table(rows=1 + len(toc_rows), cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Chapter No.", "Contents", "Page No."]
    for idx, title in enumerate(headers):
        cell = table.cell(0, idx)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title)
        set_run_font(r, size=11, bold=True)
        set_cell_shading(cell, "BFD7EE")
    for row_idx, row in enumerate(toc_rows, start=1):
        for col_idx, value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx != 1 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_run_font(r, size=11)
    add_caption(document, "Table 1: Chapter-Wise Contents of the Training Report", kind="table")
    document.add_page_break()


def add_company_profile_placeholder(document):
    add_styled_paragraph(document, "2. Company Profile", style="Heading 1", before=0, after=8)
    add_body_paragraph(
        document,
        "This section is intentionally left blank and may be completed later with the organization details after approval from the guide or institution.",
        first_line_indent=0.0,
    )
    table = document.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    placeholders = [
        ("Company / Organization Name", "........................................................"),
        ("Location", "........................................................"),
        ("Year of Establishment", "........................................................"),
        ("Domain Expertise", "........................................................"),
        ("Assigned Department / Team", "........................................................"),
    ]
    for row_idx, (label, value) in enumerate(placeholders):
        left = table.cell(row_idx, 0)
        right = table.cell(row_idx, 1)
        left.text = ""
        right.text = ""
        lp = left.paragraphs[0]
        rp = right.paragraphs[0]
        set_run_font(lp.add_run(label), size=11, bold=True)
        set_run_font(rp.add_run(value), size=11)
        set_cell_shading(left, "EAF2F8")
    add_caption(document, "Table 2: Company Profile Details to Be Filled Later", kind="table")


def add_module_table(document, member_data):
    add_styled_paragraph(document, "5. Member Module Allocation", style="Heading 1", before=12, after=8)
    add_body_paragraph(
        document,
        f"This individual training report is prepared for {member_data['member_name']} and is focused on the module areas assigned to this member during the internship and implementation period. The table below summarizes the primary training ownership and representative implementation areas considered for this report.",
        first_line_indent=0.0,
    )
    table = document.add_table(rows=1 + len(member_data["modules"]), cols=5)
    table.style = "Table Grid"
    headers = ["Sl. No.", "Module", "Frontend Areas", "Backend Areas", "Focus"]
    for idx, title in enumerate(headers):
        cell = table.cell(0, idx)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(title), size=10, bold=True)
        set_cell_shading(cell, "D9EAF7")
    for idx, module in enumerate(member_data["modules"], start=1):
        values = [
            str(idx),
            module["name"],
            ", ".join(module["frontend"]),
            ", ".join(module["backend"]),
            module["focus"],
        ]
        for col_idx, value in enumerate(values):
            cell = table.cell(idx, col_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_run_font(p.add_run(value), size=10)
    add_caption(document, "Table 3: Member-Wise Module Focus for This Training Report", kind="table")


def add_week_log(document, member_data):
    add_styled_paragraph(document, "8. Week-Wise Training Log", style="Heading 1", before=12, after=8)
    add_body_paragraph(
        document,
        "The week-wise training log is included to show how the assigned modules were studied gradually, beginning with orientation and moving toward structured implementation, testing, and workflow understanding.",
        first_line_indent=0.0,
    )
    table = document.add_table(rows=1 + len(member_data["weeks"]), cols=4)
    table.style = "Table Grid"
    headers = ["Week", "Primary Focus", "Exercise / Assignment", "Learning Outcome"]
    for idx, title in enumerate(headers):
        cell = table.cell(0, idx)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(title), size=10, bold=True)
        set_cell_shading(cell, "CFE2F3")
    for row_idx, row in enumerate(member_data["weeks"], start=1):
        values = [row["week"], row["focus"], row["task"], row["outcome"]]
        for col_idx, value in enumerate(values):
            cell = table.cell(row_idx, col_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_run_font(p.add_run(value), size=10)
    add_caption(document, "Table 6: Week-Wise Training Progress for the Assigned Modules", kind="table")
    for paragraph in member_data["week_reflections"]:
        add_body_paragraph(document, paragraph, first_line_indent=0.0)


def add_appendix(document, member_data):
    add_styled_paragraph(document, "14. Appendix", style="Heading 1", before=12, after=8)
    add_body_paragraph(
        document,
        "Only training-related screenshots and supporting material should be attached here. Project-level screenshots that do not directly represent the assigned learning modules should be avoided in this individual training report.",
        first_line_indent=0.0,
    )
    total_items = len(member_data["appendix_items"])
    for idx, description in enumerate(member_data["appendix_items"], start=1):
        box = document.add_table(rows=1, cols=1)
        box.alignment = WD_TABLE_ALIGNMENT.CENTER
        box.style = "Table Grid"
        cell = box.cell(0, 0)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run("\nInsert Training Screenshot Here\n"), size=12, bold=True, color="7F8C8D")
        p2 = cell.add_paragraph(description)
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p2.runs[0], size=10)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        add_caption(document, f"Figure {idx}: {description}", kind="figure")
        if idx % 2 == 0 and idx < total_items:
            document.add_page_break()

    add_styled_paragraph(document, "Appendix Notes", style="Heading 2", before=10, after=4)
    for note in member_data["appendix_notes"]:
        add_body_paragraph(document, note, first_line_indent=0.0)
    support_table = document.add_table(rows=1 + len(member_data["support_items"]), cols=3)
    support_table.style = "Table Grid"
    headers = ["Sl. No.", "Supporting Material", "Purpose in Training Report"]
    for idx, title in enumerate(headers):
        cell = support_table.cell(0, idx)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(title), size=10, bold=True)
        set_cell_shading(cell, "FFF2CC")
    for row_idx, item in enumerate(member_data["support_items"], start=1):
        values = [str(row_idx), item["material"], item["purpose"]]
        for col_idx, value in enumerate(values):
            cell = support_table.cell(row_idx, col_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_run_font(p.add_run(value), size=10)
    add_caption(document, "Table 9: Suggested Supporting Materials for the Appendix", kind="table")


def build_report(member_data):
    document = Document()
    apply_base_styles(document)
    configure_sections(document)
    add_cover_page(document, member_data["member_name"], member_data["report_title"])
    add_toc(document)

    add_styled_paragraph(document, "1. Introduction", style="Heading 1", before=0, after=8)
    for paragraph in member_data["introduction"]:
        add_body_paragraph(document, paragraph)

    add_company_profile_placeholder(document)

    add_styled_paragraph(document, "3. Objectives of the Internship Training", style="Heading 1", before=12, after=8)
    add_body_paragraph(document, member_data["objective_intro"], first_line_indent=0.0)
    for objective in member_data["objectives"]:
        add_bullet(document, objective)

    add_styled_paragraph(document, "4. Technology and Domain Overview", style="Heading 1", before=12, after=8)
    for paragraph in member_data["domain_overview"]:
        add_body_paragraph(document, paragraph, first_line_indent=0.0)
    for item in member_data["tech_blocks"]:
        add_styled_paragraph(document, item["title"], style="Heading 2", before=6, after=4)
        for paragraph in item["paragraphs"]:
            add_body_paragraph(document, paragraph, first_line_indent=0.0)

    add_module_table(document, member_data)

    add_styled_paragraph(document, "6. Detailed Module Study", style="Heading 1", before=12, after=8)
    add_body_paragraph(
        document,
        "This chapter provides a deeper training-oriented explanation of the assigned modules. Each subsection reflects learning activities, exercises, representative workflow understanding, and the practical observations made during the training period.",
        first_line_indent=0.0,
    )
    for index, module in enumerate(member_data["modules"], start=1):
        add_styled_paragraph(document, f"6.{index} {module['name']}", style="Heading 2", before=6, after=4)
        for paragraph in module["overview"]:
            add_body_paragraph(document, paragraph, first_line_indent=0.0)
        add_styled_paragraph(document, "Training Tasks and Exercises", style="Heading 2", before=4, after=4)
        for item in module["tasks"]:
            add_bullet(document, item)
        add_styled_paragraph(document, "Learning Outcomes", style="Heading 2", before=4, after=4)
        for paragraph in module["outcomes"]:
            add_body_paragraph(document, paragraph, first_line_indent=0.0)
        file_table = document.add_table(rows=1 + len(module["files"]), cols=3)
        file_table.style = "Table Grid"
        headers = ["Layer", "Representative Files", "Training Relevance"]
        for col_idx, title in enumerate(headers):
            cell = file_table.cell(0, col_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(p.add_run(title), size=10, bold=True)
            set_cell_shading(cell, "D9EAD3")
        for row_idx, row in enumerate(module["files"], start=1):
            values = [row["layer"], row["files"], row["relevance"]]
            for col_idx, value in enumerate(values):
                cell = file_table.cell(row_idx, col_idx)
                cell.text = ""
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                set_run_font(p.add_run(value), size=10)
        add_caption(document, f"Table 4.{index}: Representative Training References for {module['name']}", kind="table")

    add_styled_paragraph(document, "7. Module Integration and Workflow Mapping", style="Heading 1", before=12, after=8)
    for paragraph in member_data["integration"]:
        add_body_paragraph(document, paragraph, first_line_indent=0.0)
    integration_table = document.add_table(rows=1 + len(member_data["integration_map"]), cols=4)
    integration_table.style = "Table Grid"
    headers = ["Workflow Step", "Primary Module", "Connected Modules", "Reason for Integration"]
    for idx, title in enumerate(headers):
        cell = integration_table.cell(0, idx)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(title), size=10, bold=True)
        set_cell_shading(cell, "FCE5CD")
    for row_idx, row in enumerate(member_data["integration_map"], start=1):
        values = [row["step"], row["primary"], row["connected"], row["reason"]]
        for col_idx, value in enumerate(values):
            cell = integration_table.cell(row_idx, col_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_run_font(p.add_run(value), size=10)
    add_caption(document, "Table 5: Integration View of the Assigned Module Workflows", kind="table")

    add_week_log(document, member_data)

    add_styled_paragraph(document, "9. Skills Acquired", style="Heading 1", before=12, after=8)
    for section in member_data["skills"]:
        add_styled_paragraph(document, section["title"], style="Heading 2", before=4, after=4)
        for bullet in section["bullets"]:
            add_bullet(document, bullet)
        for paragraph in section.get("paragraphs", []):
            add_body_paragraph(document, paragraph, first_line_indent=0.0)

    add_styled_paragraph(document, "10. Industry Practices and Professional Exposure", style="Heading 1", before=12, after=8)
    for paragraph in member_data["industry_practices"]:
        add_body_paragraph(document, paragraph, first_line_indent=0.0)

    add_styled_paragraph(document, "11. Challenges Faced and Resolution Approach", style="Heading 1", before=12, after=8)
    for paragraph in member_data["challenges"]:
        add_body_paragraph(document, paragraph, first_line_indent=0.0)

    add_styled_paragraph(document, "12. Key Learnings and Self Assessment", style="Heading 1", before=12, after=8)
    for bullet in member_data["key_learnings"]:
        add_bullet(document, bullet)
    for paragraph in member_data["self_assessment"]:
        add_body_paragraph(document, paragraph, first_line_indent=0.0)
    assessment_table = document.add_table(rows=1 + len(member_data["assessment_table"]), cols=3)
    assessment_table.style = "Table Grid"
    headers = ["Area", "Level Achieved", "Future Improvement Focus"]
    for idx, title in enumerate(headers):
        cell = assessment_table.cell(0, idx)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(title), size=10, bold=True)
        set_cell_shading(cell, "F4CCCC")
    for row_idx, row in enumerate(member_data["assessment_table"], start=1):
        values = [row["area"], row["level"], row["future"]]
        for col_idx, value in enumerate(values):
            cell = assessment_table.cell(row_idx, col_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_run_font(p.add_run(value), size=10)
    add_caption(document, "Table 7: Self Assessment of Learning Progress", kind="table")

    add_styled_paragraph(document, "13. Conclusion", style="Heading 1", before=12, after=8)
    for paragraph in member_data["conclusion"]:
        add_body_paragraph(document, paragraph)

    add_appendix(document, member_data)

    out_path = ROOT / member_data["filename"]
    document.save(out_path)
    return out_path


MEMBERS = [
    {
        "member_name": "Darshan Jagnaure",
        "filename": "Training_Report_Darshan_Jagnaure.docx",
        "report_title": "Billing, Payment, EBL and Settlement Workflows in MERN Stack",
        "introduction": [
            "Industrial training provides practical understanding of how domain-specific software modules are designed, implemented, verified, and improved in a real development environment. This individual training report is prepared for Darshan Jagnaure and focuses on the learning journey related to billing operations, payment workflows, EBL handling, and supporting settlement modules inside a hostel management system developed using the MERN Stack.",
            "The purpose of this report is to present a training-oriented explanation of module-level work rather than a complete project description. The report therefore concentrates on practice activities such as studying the lifecycle of mess bills, observing payment capture and verification flow, understanding claim and reimbursement behavior in the EBL module, and reviewing related modules such as NOC settlement, hostel deposit adjustment, advances, and guest charges.",
            "Through these modules, the training period offered strong exposure to workflow-heavy business logic. Unlike a simple data-entry screen, billing and payment modules require careful coordination between student-facing views, caretaker actions, approval or verification flow, due amount calculation, fine handling, transaction reference tracking, and report readiness. These aspects made the module area technically rich and highly suitable for practical learning.",
            "The report also reflects how training tasks gradually moved from understanding existing workflow patterns to observing integration between frontend pages, backend controllers, data models, and financial rules. The experience helped build stronger confidence in domain logic, request-response analysis, validation, and end-to-end process thinking.",
        ],
        "objective_intro": "The training objectives for this report were centered on financial workflow modules and settlement-oriented operations in the hostel system. The main objectives were:",
        "objectives": [
            "To understand the lifecycle of hostel billing from bill generation to payment completion and status tracking.",
            "To study how payment information, UTR updates, verification flow, and fine logic are handled across student and caretaker roles.",
            "To gain practical exposure to EBL claim workflow, pending amount handling, and category-based operational tracking.",
            "To understand supporting financial modules such as advances, guest charges, NOC settlement, and hostel deposit deduction flow.",
            "To improve knowledge of how business rules are implemented in controllers, models, and frontend forms for finance-related modules.",
            "To build confidence in debugging and validating modules where multiple user actions affect a common financial record.",
        ],
        "domain_overview": [
            "Billing and payment workflows represent one of the most sensitive parts of a hostel management system because they directly affect resident dues, monthly balance visibility, and caretaker accountability. In a web-based environment, these modules must provide a clear lifecycle from amount calculation to student payment action and then to final verification and record closure.",
            "In the present system, this domain extends beyond simple mess bill display. It also includes EBL claims, extra charges, deposit adjustments, exit settlement logic, and other related records that support financial transparency during both regular hostel stay and final clearance conditions.",
        ],
        "tech_blocks": [
            {"title": "MongoDB and Financial Records", "paragraphs": ["MongoDB-based models are essential in these modules because bill, payment, EBL, NOC, deposit, and supporting settlement records require structured persistence. Through training, it became clear that data modeling plays a major role in preserving due amounts, status values, transaction details, and audit-friendly workflow states."]},
            {"title": "Express.js and Workflow Controllers", "paragraphs": ["Express controllers implement most of the financial workflow logic. Payment verification, EBL period actions, NOC calculation, deposit review, and bill-linked operations all depend on correctly structured endpoints and controller decisions. Training in this area helped develop better awareness of how business rules are encoded on the server side."]},
            {"title": "React.js for Operational Finance Screens", "paragraphs": ["The frontend portion of these modules demonstrates how finance-related actions are made understandable to users. Student screens focus on transparency and submission, whereas caretaker screens focus on review, marking, and exception handling. This role-based separation provided useful learning about interaction design for workflow systems."]},
        ],
        "modules": [
            {
                "name": "Bill Lifecycle and Payment Handling",
                "frontend": ["CaretakerBillsPage", "StudentBillsPage", "CaretakerPaymentsPage", "StudentPaymentsPage", "PaymentForm"],
                "backend": ["billRoutes.js", "messController.js", "paymentController.js", "Payment.js", "MessBill.js"],
                "focus": "Bill generation, display, payment marking, UTR flow, fine and verification logic",
                "overview": [
                    "The Bill Lifecycle and Payment Handling module was one of the most important training areas because it demonstrates how a periodic hostel due moves through multiple operational stages. These stages include bill creation, amount visibility, student-side payment intention, transaction detail update, caretaker-side confirmation, and record closure. Through this module, the training offered a complete view of how business data changes over time based on user role and payment status.",
                    "A significant learning outcome from this module was understanding that billing is not a single isolated record. It is a lifecycle. The state of a bill can change depending on whether the amount is newly generated, partially settled, pending with UTR submission, reviewed by the caretaker, or adjusted through linked logic such as fines or related settlement conditions. This lifecycle-based understanding strengthened process thinking during the training period.",
                    "The module also showed the importance of clear communication between student and caretaker views. Students need a simple interface to inspect dues and update transaction information, while caretakers need stronger controls for filtering records, validating submissions, and marking payments accurately. Building and studying both perspectives improved understanding of role-based workflow design.",
                ],
                "tasks": [
                    "Studied how mess bill records are listed and presented across student and caretaker screens.",
                    "Observed the structure of payment forms used to mark dues as paid and track UTR information.",
                    "Reviewed validation needs for payment date, payment mode, reference details, and remaining balance scenarios.",
                    "Tracked how late payment situations can trigger fine-related calculations or alerts.",
                    "Understood how student-side payment information and caretaker-side marking are kept consistent across the workflow.",
                ],
                "outcomes": [
                    "This module improved understanding of status-driven workflows and the importance of predictable transitions in financial data. It also improved awareness of where validation must happen, not only at UI level but also in backend logic.",
                    "The training around this module built practical confidence in reasoning about amount visibility, transaction references, and payment verification steps, which are critical in financial modules where even small logic errors can create trust issues.",
                ],
                "files": [
                    {"layer": "Frontend", "files": "CaretakerBillsPage.jsx, StudentBillsPage.jsx, CaretakerPaymentsPage.jsx, StudentPaymentsPage.jsx, PaymentForm.jsx", "relevance": "Used to understand display, submission, and caretaker verification flow"},
                    {"layer": "Backend", "files": "paymentController.js, messController.js, billRoutes.js", "relevance": "Used to study amount handling, status updates, and bill-payment interaction"},
                    {"layer": "Models", "files": "Payment.js, MessBill.js", "relevance": "Used to understand how bill and payment data are persisted"},
                ],
            },
            {
                "name": "EBL Workflow Module",
                "frontend": ["CaretakerEblPage", "StudentEblPage"],
                "backend": ["eblController.js", "eblRoutes.js", "EblPeriod.js", "EblCategoryReport.js"],
                "focus": "EBL claim creation, month-wise workflow, university claim handling, and remaining amount tracking",
                "overview": [
                    "The EBL Workflow Module introduced a different but equally important financial process. Unlike regular mess bill handling, EBL workflow involves a claim-oriented structure where records may be tied to categories, periods, remaining balances, and follow-up actions. This module therefore required deeper training attention on how period-driven records are prepared, reviewed, and settled over time.",
                    "A useful aspect of this module was its combination of operational complexity and user-facing clarity. Caretakers need to manage months, claim details, and selection logic, while students need visibility into what is pending, what is claimed, and what balance remains. This created strong training value because it showed how detailed backend records must still be presented through understandable UI screens.",
                    "The EBL module also strengthened understanding of exceptional billing cases. It demonstrated that not every payable amount behaves like a standard monthly bill and that systems often require dedicated modules for claims, reimbursement-linked calculations, or category-specific processing.",
                ],
                "tasks": [
                    "Studied the period-based structure of EBL records and the significance of category reporting.",
                    "Observed how caretaker actions can prepare and save university claim information.",
                    "Reviewed remaining balance scenarios and how fine alerts can apply when months are overdue.",
                    "Analyzed how student and caretaker screens must remain aligned on payable status and claim-related amounts.",
                ],
                "outcomes": [
                    "This module improved understanding of claim-driven workflows and the design of records that span more than one action cycle. It also reinforced the need for careful date-based logic and clearly explained financial summaries.",
                    "Training in this area also highlighted how a workflow module can require both operational precision and communication clarity, especially when claim and balance conditions are involved.",
                ],
                "files": [
                    {"layer": "Frontend", "files": "CaretakerEblPage.jsx, StudentEblPage.jsx", "relevance": "Used to understand the role-based EBL interaction flow"},
                    {"layer": "Backend", "files": "eblController.js, eblRoutes.js", "relevance": "Used to study period actions, claims, and balance processing"},
                    {"layer": "Models", "files": "EblPeriod.js, EblCategoryReport.js", "relevance": "Used to understand period data and reporting support"},
                ],
            },
            {
                "name": "Supporting Settlement Modules",
                "frontend": ["CaretakerAdvancesPage", "CaretakerGuestChargePage", "CaretakerNocPage", "StudentNocPage", "CaretakerHostelDepositPage", "StudentHostelDepositPage"],
                "backend": ["advanceController.js", "guestChargeController.js", "nocController.js", "hostelDepositController.js", "NocSettlement.js", "HostelDeposit.js"],
                "focus": "Advances, guest charges, NOC settlement, deposit deduction, and exit clearance support",
                "overview": [
                    "The supporting settlement modules were included in this report because they complete the broader financial workflow domain. Training exposure in this area made it clear that hostel finance is not limited to monthly dues; it also includes temporary advances, additional charges, final exit settlement, and deposit-based adjustment logic.",
                    "Among these modules, NOC settlement and hostel deposit handling were especially valuable from a training perspective because they combine amount calculation, student-provided transaction details, caretaker review, and final adjustment. The use of accepted deposit amounts in NOC flow also demonstrated how one financial module can safely interact with another when rules are clearly defined.",
                    "Studying these supporting modules broadened the understanding of settlement scenarios and showed how business processes must handle both routine and exceptional financial cases in a structured way.",
                ],
                "tasks": [
                    "Reviewed how advances and guest charges can be tracked as supporting financial records.",
                    "Studied the workflow where hostel deposits are submitted, reviewed, and later used for NOC deduction.",
                    "Observed how NOC records combine base amount, damages, others, total balance, and deposit usage.",
                    "Understood the role of caretaker review in confirming the final amount before exit-related settlement is closed.",
                ],
                "outcomes": [
                    "This module area improved understanding of settlement completeness and showed how a mature hostel system supports both monthly and exit-oriented financial workflows.",
                    "The training also strengthened confidence in reasoning about cross-module dependencies, especially where one accepted financial record affects another final settlement record.",
                ],
                "files": [
                    {"layer": "Frontend", "files": "CaretakerNocPage.jsx, StudentNocPage.jsx, CaretakerHostelDepositPage.jsx, StudentHostelDepositPage.jsx", "relevance": "Used to understand settlement and deposit deduction interactions"},
                    {"layer": "Backend", "files": "nocController.js, hostelDepositController.js, advanceController.js, guestChargeController.js", "relevance": "Used to study support workflows and final settlement logic"},
                    {"layer": "Models", "files": "NocSettlement.js, HostelDeposit.js, Advance.js, GuestCharge.js", "relevance": "Used to understand structured storage of settlement-related records"},
                ],
            },
        ],
        "integration": [
            "The assigned modules for this report are strongly integrated because a bill or EBL amount often moves through the same broad lifecycle: record creation, amount visibility, submission or payment, caretaker review, and final closure or adjustment. This similarity made the training particularly useful because it allowed repeated study of workflow design across different financial cases.",
            "Another important integration point is the relationship between routine payment handling and final settlement support. NOC and hostel deposit modules are not monthly billing modules, but they still depend on the same principles of amount clarity, student input, review flow, and safe record updates. Observing these connections helped strengthen system-level understanding.",
        ],
        "integration_map": [
            {"step": "Monthly due visibility", "primary": "Bill Lifecycle", "connected": "Payment Handling", "reason": "Bill records become meaningful only when linked to student-facing due and payment actions"},
            {"step": "Claim-driven payable flow", "primary": "EBL Workflow", "connected": "Payment Handling", "reason": "Remaining amount and overdue handling must eventually connect to payment or settlement actions"},
            {"step": "Exit settlement preparation", "primary": "NOC Settlement", "connected": "Hostel Deposit", "reason": "Accepted deposit can reduce payable balance during clearance"},
            {"step": "Additional amount support", "primary": "Guest Charges / Advances", "connected": "Bill and Payment workflows", "reason": "Supporting financial records may influence the broader settlement view"},
        ],
        "weeks": [
            {"week": "1", "focus": "Bill workflow orientation", "task": "Studied bill listing, due presentation, and screen flow", "outcome": "Understood lifecycle view of monthly dues"},
            {"week": "2", "focus": "Payment submission and marking", "task": "Reviewed payment forms, UTR input, and verification points", "outcome": "Improved understanding of student-to-caretaker payment workflow"},
            {"week": "3", "focus": "Fine and due-state scenarios", "task": "Observed late-payment conditions and payable state changes", "outcome": "Learned why financial status must be explicit and predictable"},
            {"week": "4", "focus": "EBL module study", "task": "Analyzed period records, claims, and balance handling", "outcome": "Understood claim-oriented financial workflows"},
            {"week": "5", "focus": "Supporting financial records", "task": "Reviewed advances and guest charge concepts", "outcome": "Expanded knowledge beyond standard monthly billing"},
            {"week": "6", "focus": "NOC and hostel deposit logic", "task": "Studied exit settlement and deposit deduction workflow", "outcome": "Understood cross-module financial dependency"},
            {"week": "7", "focus": "Workflow integration review", "task": "Mapped how payment and settlement modules interact", "outcome": "Improved system-level process reasoning"},
            {"week": "8", "focus": "Validation and reporting awareness", "task": "Reviewed how accurate financial handling supports reporting", "outcome": "Connected operational modules with record reliability"},
        ],
        "week_reflections": [
            "The week-wise progression helped convert what initially seemed like multiple unrelated finance screens into a connected operational system. Each week added a new layer of understanding, beginning from bills and ending with final settlement-oriented processes.",
            "A major reflection from this timeline is that financial modules are easier to understand when tracked through state changes rather than just screen names. The training therefore improved not only code-related understanding but also process-oriented thinking.",
        ],
        "skills": [
            {"title": "Technical Skills", "bullets": ["Better understanding of bill lifecycle and amount status transitions", "Improved knowledge of payment verification, UTR handling, and fine-aware workflow logic", "Awareness of EBL, NOC, hostel deposit, and settlement-linked financial modules"], "paragraphs": ["The technical skill growth in this report area was centered on workflow consistency, amount correctness, and cross-module reasoning. These are critical skills for domains involving financial trust and operational accountability."]},
            {"title": "Tools and Software", "bullets": ["VS Code for studying module structure and interaction flow", "Postman-style API understanding for request and response thinking", "MongoDB model awareness for financial persistence"], "paragraphs": ["The tools in these modules were useful not merely for coding but for verifying how financial logic travels between interfaces, endpoints, and stored records."]},
            {"title": "Soft Skills", "bullets": ["Improved patience while analyzing business-rule-heavy modules", "Better communication around workflow steps and financial state logic", "Stronger documentation habit for multi-step operational records"], "paragraphs": []},
        ],
        "industry_practices": [
            "This module group provided practical exposure to an important industry principle: workflow clarity is as important as feature completeness. In a finance-related domain, users must know what stage a record is in, who must act next, and whether a payment or adjustment has been confirmed.",
            "Another important observation was that domain rules often span multiple modules. A payment record cannot always be understood in isolation if it is influenced by claim flow, fine conditions, or deposit deductions. Training in this area therefore strengthened appreciation for careful integration design.",
        ],
        "challenges": [
            "One challenge in this module area was understanding the variety of financial cases handled by the system. Monthly bills, claim-oriented EBL flow, deposit adjustment, and NOC settlement each behave differently even though they all relate to money movement.",
            "Another challenge was following role-based sequence correctly. Student actions and caretaker actions do not happen in the same order across every module, so learning the exact lifecycle required repeated review of transitions and state changes.",
            "These challenges were addressed by studying each module separately first, then mapping the overlaps between them through workflow comparison and representative file review.",
        ],
        "key_learnings": [
            "Financial modules are best understood as state-driven workflows rather than as isolated data forms.",
            "Student-facing simplicity often depends on complex caretaker-side review and backend validation logic.",
            "Settlement modules such as NOC and deposit deduction strengthen the completeness of the financial domain.",
            "Fine handling and overdue awareness are important for operational transparency.",
        ],
        "self_assessment": [
            "The training period improved confidence in understanding workflow-heavy financial modules and how they connect across routine and exceptional hostel scenarios. It also improved the ability to distinguish between display logic, workflow status logic, and settlement logic.",
            "Further improvement is still needed in advanced financial reporting, broader reconciliation logic, and long-term audit-oriented workflow modeling. However, the current training has created a strong conceptual base for those future steps.",
        ],
        "assessment_table": [
            {"area": "Bill workflow understanding", "level": "Strong working understanding", "future": "Study more advanced billing automation"},
            {"area": "Payment verification logic", "level": "Comfortable with practical flow", "future": "Improve exception and retry handling knowledge"},
            {"area": "EBL process reasoning", "level": "Good workflow awareness", "future": "Explore deeper reporting and aggregation"},
            {"area": "Settlement modules", "level": "Good conceptual clarity", "future": "Improve audit and reconciliation perspective"},
        ],
        "conclusion": [
            "The training covered in this report provided strong exposure to financial workflows inside a hostel management environment. Through bill handling, payment processing, EBL claim flow, and settlement-related modules, the training period created a deeper understanding of how real operational systems manage payable records from creation to closure.",
            "This report area also highlighted the importance of accuracy, clarity, and workflow discipline in software modules that influence user trust. As a result, the training not only improved technical understanding but also strengthened process awareness and responsibility toward data-sensitive module design.",
        ],
        "appendix_items": [
            "Screenshot of mess bill listing or bill lifecycle exercise",
            "Screenshot of payment form or UTR update practice",
            "Screenshot of caretaker payment marking screen",
            "Screenshot of student payment visibility module",
            "Screenshot of overdue or fine-related training observation",
            "Screenshot of EBL month selection or claim preparation",
            "Screenshot of EBL balance or category report training screen",
            "Screenshot of NOC amount split and settlement exercise",
            "Screenshot of hostel deposit acceptance or review flow",
            "Screenshot of deposit deduction checkbox in NOC training flow",
            "Screenshot of advances or guest charge practice module",
            "Screenshot of backend route/controller study notes for billing",
            "Screenshot of payment status debugging or validation exercise",
            "Screenshot of settlement workflow mapping notes",
            "Screenshot of module integration checklist for financial workflows",
            "Screenshot of build or test verification related to billing modules",
            "Screenshot of training notes for EBL workflow understanding",
            "Screenshot of role-based sequence mapping for payment modules",
        ],
        "appendix_notes": [
            "Screenshots must correspond only to the training activities for billing, payment, EBL, and related settlement modules.",
            "The appendix should avoid unrelated project screens unless they directly demonstrate the assigned training module workflow.",
            "Each screenshot may be labeled with a short note describing the practical exercise performed during training.",
        ],
        "support_items": [
            {"material": "Bill screen screenshot", "purpose": "Shows monthly due visibility and bill lifecycle understanding"},
            {"material": "Payment submission screenshot", "purpose": "Shows transaction update and payment workflow learning"},
            {"material": "EBL workflow screenshot", "purpose": "Represents claim and balance training activities"},
            {"material": "NOC/deposit screenshot", "purpose": "Shows settlement and deduction understanding"},
            {"material": "Controller study notes", "purpose": "Documents backend workflow analysis"},
        ],
    },
    {
        "member_name": "Vaishnavi",
        "filename": "Training_Report_Vaishnavi.docx",
        "report_title": "Expense, Hostel Expense, Consumption and Charge Management in MERN Stack",
        "introduction": [
            "This individual training report is prepared for Vaishnavi and focuses on the training modules related to expense records, hostel expense tracking, consumption management, and charge calculation workflows in the hostel management system. These modules are operationally important because they influence cost distribution, transparency of shared usage, and structured record keeping for hostel administration.",
            "The training area covered in this report differs from general project-level description because it concentrates on domain-specific learning exercises such as adding and managing expenses, understanding hostel-level expenditure flow, recording consumption data, and studying how charge records are calculated and presented. The module combination provided strong exposure to both data management and operational accounting logic.",
            "A key value of these modules is that they sit at the intersection of administration, allocation, and accountability. Consumption records and charges affect how shared resources are interpreted, while expense and hostel expense modules reflect how costs are documented and prepared for broader reporting or review. This made the training especially valuable for understanding the operational backbone of the system.",
            "The report therefore presents a detailed training-oriented discussion of these modules, the exercises performed around them, the integration points observed during the learning period, and the professional insights developed while working in this domain area.",
        ],
        "objective_intro": "The main training objectives for this report were focused on operational cost tracking and resource-based record management. The objectives were:",
        "objectives": [
            "To understand how regular expenses and hostel-specific expenses are recorded, organized, and reviewed.",
            "To study consumption tracking workflow and its importance for fair and transparent charge distribution.",
            "To gain practical exposure to charge management screens, input structures, and related backend logic.",
            "To improve understanding of how resource usage data is linked to cost management processes.",
            "To study how operational records support later reporting, analysis, and administrative decision making.",
            "To strengthen understanding of role-based screens used by caretakers and administrators for cost-related modules.",
        ],
        "domain_overview": [
            "Expense, consumption, and charge management modules form the administrative calculation layer of the system. They do not merely store records; they provide the structured inputs needed to understand how hostel resources are used, how costs are incurred, and how financial responsibility can be distributed or reviewed with clarity.",
            "From a training perspective, these modules are valuable because they combine data-entry discipline, calculated reasoning, and operational transparency. They show that a hostel management system is not only about student records and payments, but also about maintaining accurate internal operational cost records that support responsible management.",
        ],
        "tech_blocks": [
            {"title": "MongoDB and Operational Data", "paragraphs": ["Expense and consumption workflows rely on consistent storage of recurring operational records. Training in this area improved appreciation for clear schema design, meaningful field structure, and reliable categorization of records related to usage and cost."]},
            {"title": "Express.js for Record Management", "paragraphs": ["Controllers for expense, hostel expense, consumption, and charges demonstrate how server-side logic supports filtering, saving, reviewing, and updating structured records. These controllers provided good training exposure to administrative workflow implementation."]},
            {"title": "React.js for Administrative Input Screens", "paragraphs": ["These modules rely heavily on caretaker and admin input screens. React-based forms, management pages, and calculation helpers were useful for understanding how a structured operational workflow is translated into practical user interfaces."]},
        ],
        "modules": [
            {
                "name": "Expense and Hostel Expense Management",
                "frontend": ["AdminExpensesPage", "CaretakerExpensesPage", "ExpenseManagementPage", "CaretakerHostelExpensePage"],
                "backend": ["expenseController.js", "hostelExpenseController.js", "expenseRoutes.js", "hostelExpenseRoutes.js", "Expense.js", "HostelExpense.js"],
                "focus": "Expense entry, hostel expense tracking, categorization, and operational record maintenance",
                "overview": [
                    "Expense and Hostel Expense Management was a strong training area because it highlighted the importance of structured administrative records in a hostel environment. Through these modules, it became clear that operational expenditure must be captured with consistent fields, usable categories, and traceable context so that later reporting and review remain reliable.",
                    "Training in this module also showed that there can be more than one layer of expense handling. General expenses and hostel-specific expenses may share similar record behavior, but their operational role can differ. Observing both helped in understanding how domain boundaries are reflected in module design.",
                    "Another useful learning point from this module was the relationship between clean input forms and later data quality. If operational expenses are not recorded carefully at entry stage, downstream summaries and decision support can become confusing. This made the module very effective for understanding the value of disciplined data capture.",
                ],
                "tasks": [
                    "Studied form structures for recording expense details and hostel expense records.",
                    "Reviewed how categories, dates, values, and descriptions support clear operational bookkeeping.",
                    "Observed differences between general expense flow and hostel-level expense flow.",
                    "Analyzed how these records support later management review and reporting.",
                ],
                "outcomes": [
                    "This module improved understanding of administrative record discipline and the relationship between input quality and reporting quality.",
                    "It also strengthened awareness of how operational cost modules support accountability within hostel administration.",
                ],
                "files": [
                    {"layer": "Frontend", "files": "AdminExpensesPage.jsx, CaretakerExpensesPage.jsx, ExpenseManagementPage.jsx, CaretakerHostelExpensePage.jsx", "relevance": "Used to understand caretaker/admin workflows for expense records"},
                    {"layer": "Backend", "files": "expenseController.js, hostelExpenseController.js, expenseRoutes.js, hostelExpenseRoutes.js", "relevance": "Used to study saving, listing, and handling expense records"},
                    {"layer": "Models", "files": "Expense.js, HostelExpense.js", "relevance": "Used to understand structured storage of operational costs"},
                ],
            },
            {
                "name": "Consumption and Charge Management",
                "frontend": ["AdminConsumptionPage", "CaretakerConsumptionPage", "ConsumptionManagementPage", "CaretakerChargesPage", "ConsumptionForm", "ChargeForm"],
                "backend": ["consumptionController.js", "chargeController.js", "consumptionRoutes.js", "chargeRoutes.js", "StudentConsumption.js", "Charge.js"],
                "focus": "Consumption tracking, charge records, usage interpretation, and cost distribution logic",
                "overview": [
                    "Consumption and Charge Management formed the second major training domain for this report. These modules are operationally important because they capture usage-oriented data and transform it into charge-related understanding. This made the training interesting because the records are not static administrative entries; they often represent measured or derived values that later influence financial interpretation.",
                    "A key learning point from this module was the importance of fairness and clarity in resource-based workflows. Consumption data must be handled carefully because it can influence how users understand shared costs or hostel resource allocation. The system therefore requires structured input, understandable summaries, and reliable storage.",
                    "The module also provided useful exposure to calculation support and helper logic. Compared to simple CRUD screens, charge and consumption modules demonstrate a slightly more interpretive workflow where raw values may later support allocation, review, or explanation within the hostel system.",
                ],
                "tasks": [
                    "Studied how consumption data is entered, organized, and reviewed.",
                    "Observed how charge-related records are created and maintained from usage-oriented inputs.",
                    "Reviewed the helper and form logic used for calculations and operational clarity.",
                    "Analyzed how consumption and charge records may support hostel-level cost understanding.",
                ],
                "outcomes": [
                    "This module improved understanding of operational data that connects usage with financial interpretation.",
                    "It also built stronger confidence in studying records that require both input structure and calculation awareness.",
                ],
                "files": [
                    {"layer": "Frontend", "files": "AdminConsumptionPage.jsx, CaretakerConsumptionPage.jsx, ConsumptionManagementPage.jsx, CaretakerChargesPage.jsx, ConsumptionForm.jsx, ChargeForm.jsx", "relevance": "Used to study administrative handling of usage and charge records"},
                    {"layer": "Backend", "files": "consumptionController.js, chargeController.js, consumptionRoutes.js, chargeRoutes.js", "relevance": "Used to understand save, update, and calculation support workflows"},
                    {"layer": "Models", "files": "StudentConsumption.js, Charge.js", "relevance": "Used to understand persistence of usage and charge-related records"},
                ],
            },
        ],
        "integration": [
            "The expense, hostel expense, consumption, and charge modules are strongly connected because they all contribute to the operational cost understanding of the hostel. While not all of them directly present student-facing payment records, they create the structured administrative data that supports accurate management, review, and later reporting.",
            "Consumption records and charge modules are especially important because they create a bridge between usage and cost. Expense modules, on the other hand, strengthen the administrative side of bookkeeping. Studying them together created a broader understanding of how internal hostel operations influence system-level financial clarity.",
        ],
        "integration_map": [
            {"step": "Expense entry", "primary": "Expense Management", "connected": "Hostel Expense Management", "reason": "Both maintain categorized operational cost records"},
            {"step": "Usage capture", "primary": "Consumption Management", "connected": "Charge Management", "reason": "Consumption understanding supports charge interpretation"},
            {"step": "Administrative review", "primary": "Expense / Charge modules", "connected": "Reporting and Analytics", "reason": "Operational records later support summaries and decision-making"},
            {"step": "Cost clarity", "primary": "Hostel Expense and Charges", "connected": "Broader finance workflow", "reason": "Organized cost records strengthen transparency across the system"},
        ],
        "weeks": [
            {"week": "1", "focus": "Expense module orientation", "task": "Studied general expense record structure", "outcome": "Understood administrative cost entry basics"},
            {"week": "2", "focus": "Hostel expense flow", "task": "Reviewed hostel-specific expense screens and records", "outcome": "Learned domain-specific expenditure handling"},
            {"week": "3", "focus": "Consumption record study", "task": "Observed usage entry and record organization", "outcome": "Understood operational resource tracking"},
            {"week": "4", "focus": "Charge management", "task": "Reviewed charge forms and related calculation support", "outcome": "Learned how usage can connect to charge records"},
            {"week": "5", "focus": "Form and helper logic", "task": "Studied administrative input discipline and calculation helpers", "outcome": "Improved understanding of operational input quality"},
            {"week": "6", "focus": "Cross-module review", "task": "Mapped expense, consumption, and charge relationships", "outcome": "Improved system-level cost workflow understanding"},
            {"week": "7", "focus": "Operational review mindset", "task": "Observed how record quality affects later reporting", "outcome": "Connected input discipline with transparency"},
            {"week": "8", "focus": "Summary and validation", "task": "Reviewed module responsibilities and role-based usage", "outcome": "Built stronger confidence in cost-management modules"},
        ],
        "week_reflections": [
            "The week-wise training flow helped establish a useful distinction between cost-entry modules and usage-linked modules. This distinction improved understanding of how different record types can still support the same broader administrative objective.",
            "A strong outcome from this sequence was the realization that well-organized operational records are necessary long before report generation or analytics begins. Clean inputs create better summaries and better administrative trust.",
        ],
        "skills": [
            {"title": "Technical Skills", "bullets": ["Improved understanding of operational expense record handling", "Better awareness of consumption and charge workflow structure", "Greater confidence in studying calculation-supporting administrative modules"], "paragraphs": ["This training area strengthened understanding of modules where accurate input and categorized storage are essential for later analysis and transparency."]},
            {"title": "Tools and Software", "bullets": ["React forms and management pages for operational entry", "Backend controller awareness for record handling", "MongoDB models for persistent cost and usage records"], "paragraphs": ["The tools and code structure in these modules made it easier to understand how administrative workflows remain organized over time."]},
            {"title": "Soft Skills", "bullets": ["Attention to detail during record interpretation", "Patience in studying administrative workflow patterns", "Improved communication of operational logic"], "paragraphs": []},
        ],
        "industry_practices": [
            "These modules highlighted an important industry practice: cost-related systems depend heavily on disciplined data capture. Strong analysis and reporting are impossible without good operational records.",
            "The training also showed that usage-based modules often require thoughtful explanation because numbers alone are not always meaningful without context. This is why good UI structure, labeling, and backend consistency are important in cost-management software.",
        ],
        "challenges": [
            "One challenge was understanding the difference between various cost-related records and why separate modules were needed for expenses, hostel expenses, consumption, and charges.",
            "Another challenge was interpreting how usage-oriented records can influence later cost understanding without always directly appearing as immediate student-facing bills.",
            "These challenges were reduced by reviewing the workflow role of each module separately and then comparing how all of them contribute to operational transparency.",
        ],
        "key_learnings": [
            "Operational records become more useful when categories, descriptions, and values are captured consistently.",
            "Consumption and charge workflows require both data accuracy and interpretive clarity.",
            "Administrative modules are foundational for later reporting and financial transparency.",
            "Good input design directly improves data quality and downstream usefulness.",
        ],
        "self_assessment": [
            "The training in this report area created stronger confidence in understanding operational modules that support internal hostel cost management. It improved the ability to study systems where not every record leads immediately to payment, but still plays a crucial role in administrative accuracy.",
            "Future improvement is still needed in deeper cost analysis logic, advanced formula support, and stronger report-driven interpretation of consumption and charge patterns. However, the current training has created a reliable practical foundation in these modules.",
        ],
        "assessment_table": [
            {"area": "Expense workflow understanding", "level": "Good working understanding", "future": "Study advanced categorization and audit patterns"},
            {"area": "Hostel expense handling", "level": "Comfortable with administrative flow", "future": "Improve comparative and periodic review perspective"},
            {"area": "Consumption logic", "level": "Good conceptual clarity", "future": "Explore more advanced allocation logic"},
            {"area": "Charge module awareness", "level": "Comfortable with form and data flow", "future": "Improve calculation-depth understanding"},
        ],
        "conclusion": [
            "The training modules documented in this report provided meaningful practical exposure to the administrative and operational cost side of the hostel management system. Through expense, hostel expense, consumption, and charge management, the training period strengthened understanding of how internal records support transparent management.",
            "The overall experience helped develop a more mature view of data discipline, operational clarity, and the role of well-maintained records in building trustworthy software systems. These outcomes make the training highly relevant for future work in administrative and process-driven applications.",
        ],
        "appendix_items": [
            "Screenshot of expense entry or expense management exercise",
            "Screenshot of hostel expense training page or form",
            "Screenshot of expense category or record review activity",
            "Screenshot of consumption record entry exercise",
            "Screenshot of consumption management screen",
            "Screenshot of charge form or charge calculation training activity",
            "Screenshot of caretaker charge management page",
            "Screenshot of admin consumption view or data review",
            "Screenshot of helper calculation notes for hostel expense logic",
            "Screenshot of record-edit or update workflow exercise",
            "Screenshot of operational cost validation practice",
            "Screenshot of backend controller study notes for expense modules",
            "Screenshot of model or schema review related to consumption records",
            "Screenshot of module relationship notes for cost workflows",
            "Screenshot of test or verification activity for expense records",
            "Screenshot of training checklist for consumption and charges",
            "Screenshot of usage-to-charge mapping notes",
            "Screenshot of role-based cost review workflow observation",
        ],
        "appendix_notes": [
            "Use only screenshots that directly relate to expense, hostel expense, consumption, or charge management training activities.",
            "Avoid screenshots from unrelated modules unless they clearly demonstrate the assigned training context.",
            "Each screenshot can include a short note on what was practiced and what was understood from that activity.",
        ],
        "support_items": [
            {"material": "Expense module screenshot", "purpose": "Shows administrative cost-record workflow"},
            {"material": "Consumption module screenshot", "purpose": "Represents resource usage training activity"},
            {"material": "Charge module screenshot", "purpose": "Shows how usage links to charge understanding"},
            {"material": "Schema/controller notes", "purpose": "Documents backend training observations"},
            {"material": "Verification checklist", "purpose": "Shows structured learning and review"},
        ],
    },
    {
        "member_name": "Pooja",
        "filename": "Training_Report_Pooja.docx",
        "report_title": "Reporting, Approval, Analytics and Notification Modules in MERN Stack",
        "introduction": [
            "This individual training report is prepared for Pooja and focuses on the training journey related to reporting workflows, approval-oriented screens, analytics views, and notification mechanisms in the hostel management system. These modules are especially important because they help convert operational records into decision-support information and communication actions.",
            "Unlike data-entry-heavy modules, the modules covered in this report emphasize review, oversight, summary visibility, and information delivery. This made the training highly useful because it introduced the idea that mature software systems require more than transaction screens; they also require modules that help administrators observe trends, approve requests, and communicate effectively.",
            "The report concentrates on training activities such as studying report generation and action panels, understanding approval responsibilities on admin screens, observing analytics-oriented summaries, and reviewing notification handling as a communication support module. This training focus created a broad perspective on supervisory workflows inside the system.",
            "The following chapters therefore document how these modules were understood, what exercises and analysis were performed around them, and what professional insights emerged from studying modules that sit closer to decision support and oversight.",
        ],
        "objective_intro": "The training objectives for this report focused on oversight, summary, and communication-oriented modules. The main objectives were:",
        "objectives": [
            "To understand how reporting modules organize and present structured hostel records for review and action.",
            "To study admin approval workflows and the importance of controlled decision points in the system.",
            "To gain practical exposure to analytics modules that summarize trends, counts, and operational performance indicators.",
            "To understand how notification handling supports communication and event awareness across the system.",
            "To improve awareness of how supervisory modules depend on clean data coming from other operational areas.",
            "To build confidence in studying role-based interfaces designed for review, analysis, and communication.",
        ],
        "domain_overview": [
            "Reporting, approval, analytics, and notification modules are different from routine input modules because they operate at a higher level of system maturity. They rely on information collected from multiple workflows and transform that information into summaries, decision points, or communication outputs that help users act with better awareness.",
            "From a training perspective, these modules are valuable because they encourage systems thinking. They show how the value of a software platform grows when raw data is not only stored but also reviewed, interpreted, approved, and communicated in a structured way.",
        ],
        "tech_blocks": [
            {"title": "MongoDB and Summary-Ready Data", "paragraphs": ["These modules depend on reliable data across the system. Training in this area improved awareness that reports and analytics are only as trustworthy as the underlying records stored through models and operational workflows."]},
            {"title": "Express.js for Supervisory Workflows", "paragraphs": ["Controllers related to reports, analytics, and notifications demonstrate how server-side logic can aggregate, approve, prepare, or deliver information rather than simply storing it. This broadened understanding of backend responsibilities."]},
            {"title": "React.js for Review and Decision Support Screens", "paragraphs": ["Admin and caretaker reporting screens, analytics pages, and approval cards provided useful examples of how UI can support supervision, not just data entry. This made the training especially relevant for understanding higher-level application behavior."]},
        ],
        "modules": [
            {
                "name": "Reporting and Approval Module",
                "frontend": ["CaretakerReportsPage", "ReportActionPanel", "AdminApprovalsPage", "ApprovalCard"],
                "backend": ["reportController.js", "monthlyExpenseReportController.js", "pdfController.js", "pdfDocumentController.js", "adminController.js", "Report.js", "MonthlyExpenseReport.js"],
                "focus": "Report generation, review actions, approval handling, and administrative oversight",
                "overview": [
                    "The Reporting and Approval Module was a highly meaningful training area because it showed how the system supports supervision and formal review. Reports are important not only for record keeping but also for converting day-to-day operational data into an understandable format for caretaker and admin decision making.",
                    "Approval workflows added another layer of learning by demonstrating controlled review points inside the system. These modules help ensure that not every action becomes final without proper oversight. This made the training valuable for understanding accountability and the role of supervised operations in administrative software.",
                    "The module also highlighted how reporting screens must balance detail with usability. A good report or action panel should provide enough information for decision making without creating confusion. Observing this balance provided useful interface-level learning alongside backend logic understanding.",
                ],
                "tasks": [
                    "Studied report-oriented pages and action panels used for review workflows.",
                    "Observed approval cards and admin approval screens to understand supervisory actions.",
                    "Reviewed how PDF and monthly reporting controllers may support export or structured summary generation.",
                    "Analyzed how report quality depends on consistency of source records from other modules.",
                ],
                "outcomes": [
                    "This module improved understanding of supervisory software behavior and the importance of controlled actions before final decisions or report use.",
                    "It also strengthened awareness of how structured summaries can improve operational clarity across the system.",
                ],
                "files": [
                    {"layer": "Frontend", "files": "CaretakerReportsPage.jsx, ReportActionPanel.jsx, AdminApprovalsPage.jsx, ApprovalCard.jsx", "relevance": "Used to study review, report interaction, and approval-related UI"},
                    {"layer": "Backend", "files": "reportController.js, monthlyExpenseReportController.js, pdfController.js, pdfDocumentController.js, adminController.js", "relevance": "Used to understand reporting logic, document support, and approval handling"},
                    {"layer": "Models", "files": "Report.js, MonthlyExpenseReport.js", "relevance": "Used to understand persistence and structure of report records"},
                ],
            },
            {
                "name": "Analytics and Notification Module",
                "frontend": ["AdminAnalyticsPage"],
                "backend": ["analyticsController.js", "notificationController.js", "notificationRoutes.js", "Notification.js"],
                "focus": "Analytics visibility, summary interpretation, event awareness, and communication support",
                "overview": [
                    "The Analytics and Notification Module contributed a different type of training exposure by focusing on awareness rather than direct transaction entry. Analytics pages help users understand patterns, counts, or high-level system activity, while notification modules support timely communication and event awareness.",
                    "This module was useful because it highlighted the idea that software systems should not only store records but also help users notice important conditions. Notifications support responsiveness, and analytics support interpretation. Together they create a more intelligent and usable administrative platform.",
                    "The training also showed that these modules depend heavily on good upstream data quality. Weak or inconsistent source records can make analytics less meaningful and notifications less helpful. This reinforced the system-wide importance of reliable operational modules.",
                ],
                "tasks": [
                    "Studied analytics pages and summary-oriented views for administrative observation.",
                    "Reviewed how notification records and notification handling support communication flow.",
                    "Analyzed how high-level views depend on clean source data from reports, expenses, billing, and other modules.",
                    "Observed how decision support and communication modules improve system usability at a higher level.",
                ],
                "outcomes": [
                    "This module improved understanding of summary-oriented interfaces and event-driven support behavior.",
                    "It also strengthened appreciation for how dashboards and notifications increase the practical usefulness of a management platform.",
                ],
                "files": [
                    {"layer": "Frontend", "files": "AdminAnalyticsPage.jsx", "relevance": "Used to understand administrative analytics and summary presentation"},
                    {"layer": "Backend", "files": "analyticsController.js, notificationController.js, notificationRoutes.js", "relevance": "Used to study summary preparation and communication handling"},
                    {"layer": "Models", "files": "Notification.js", "relevance": "Used to understand storage of communication-support records"},
                ],
            },
        ],
        "integration": [
            "The modules covered in this report show the upper supervisory layer of the system. Reporting and approvals depend on operational data quality, analytics depend on structured aggregation, and notifications support timely awareness of events or required attention points.",
            "Studying these modules together created a strong understanding of how review, interpretation, and communication can make a system more complete. They are not isolated utilities; they are the modules that help administrators act intelligently on the basis of system information.",
        ],
        "integration_map": [
            {"step": "Operational summary generation", "primary": "Reporting", "connected": "Expense, payment, and other records", "reason": "Reports require accurate source module data"},
            {"step": "Decision point handling", "primary": "Approvals", "connected": "Admin oversight", "reason": "Approval screens convert record review into controlled actions"},
            {"step": "Trend visibility", "primary": "Analytics", "connected": "Reporting and operational data", "reason": "Analytics help interpret system performance or status at a higher level"},
            {"step": "Event communication", "primary": "Notifications", "connected": "Workflow events and user awareness", "reason": "Notifications improve responsiveness and usability"},
        ],
        "weeks": [
            {"week": "1", "focus": "Report module orientation", "task": "Studied report pages and summary usage", "outcome": "Understood supervisory record presentation"},
            {"week": "2", "focus": "Approval flow review", "task": "Observed approval cards and admin actions", "outcome": "Learned controlled review workflow"},
            {"week": "3", "focus": "Document support awareness", "task": "Reviewed PDF/report-support controllers", "outcome": "Expanded understanding of report output readiness"},
            {"week": "4", "focus": "Analytics module study", "task": "Observed summary dashboards and count-based visibility", "outcome": "Understood higher-level administrative insight"},
            {"week": "5", "focus": "Notification module understanding", "task": "Reviewed event awareness and notification behavior", "outcome": "Learned communication-support workflow"},
            {"week": "6", "focus": "Cross-module dependency review", "task": "Mapped how reports depend on source record quality", "outcome": "Improved system-level thinking"},
            {"week": "7", "focus": "Decision support perspective", "task": "Studied relation between analytics, approvals, and reports", "outcome": "Understood supervisory layer integration"},
            {"week": "8", "focus": "Summary and reflection", "task": "Reviewed module responsibilities and future learning areas", "outcome": "Strengthened confidence in oversight-oriented modules"},
        ],
        "week_reflections": [
            "The week-wise training sequence helped reveal the supervisory nature of these modules. Unlike raw transaction screens, these modules become meaningful when viewed in relation to the broader system data they depend on.",
            "A major insight from the weekly progression was that high-level modules add practical value only when the lower-level operational modules are already structured and reliable. This deepened appreciation for system-wide data quality.",
        ],
        "skills": [
            {"title": "Technical Skills", "bullets": ["Improved understanding of report and approval workflow structure", "Better awareness of analytics-oriented summary screens", "Practical understanding of communication-support modules like notifications"], "paragraphs": ["This training area broadened technical understanding beyond CRUD workflows and introduced the importance of aggregation, interpretation, and oversight."]},
            {"title": "Tools and Software", "bullets": ["Admin and caretaker report interfaces", "Backend controllers for report and notification flow", "Document-support awareness through PDF-related controllers"], "paragraphs": []},
            {"title": "Soft Skills", "bullets": ["Analytical thinking while interpreting summaries", "Decision-awareness in reviewing approval workflows", "Improved communication of module relationships"], "paragraphs": []},
        ],
        "industry_practices": [
            "The training in these modules highlighted a key industry practice: systems become more useful when raw operations are transformed into meaningful summaries and timely communication. Reports, approvals, analytics, and notifications are the modules that often make management systems truly practical.",
            "Another important observation was that supervisory modules demand clarity and trust. If an approval flow is confusing or a report summary is unreliable, decision support weakens. This makes consistency and traceability particularly important in such modules.",
        ],
        "challenges": [
            "One challenge was understanding modules that are less transaction-oriented and more summary-oriented, because their value depends on information coming from many other parts of the system.",
            "Another challenge was tracing how notifications, analytics, and report views connect conceptually even when they appear in different screens or controllers.",
            "These challenges were managed by studying each module as a role-based support layer first and then mapping the dependencies between them.",
        ],
        "key_learnings": [
            "Reports and approvals create accountability and structured administrative oversight.",
            "Analytics modules improve visibility, but only when supported by reliable source data.",
            "Notifications increase system responsiveness by improving event awareness.",
            "Supervisory modules represent a higher layer of maturity in software system design.",
        ],
        "self_assessment": [
            "The training recorded in this report improved confidence in understanding modules that present, summarize, and communicate information rather than simply storing it. This was an important learning expansion because such modules are central to real administrative decision support.",
            "Future learning is still needed in advanced dashboard design, report export optimization, and richer notification strategies. However, the current training has already created strong conceptual clarity in these important oversight-related modules.",
        ],
        "assessment_table": [
            {"area": "Report module understanding", "level": "Good practical clarity", "future": "Study richer export and summarization patterns"},
            {"area": "Approval workflow awareness", "level": "Comfortable with review logic", "future": "Explore more complex approval chains"},
            {"area": "Analytics interpretation", "level": "Good conceptual understanding", "future": "Improve metric selection and dashboard design"},
            {"area": "Notification module knowledge", "level": "Working understanding", "future": "Study event-driven notification strategies"},
        ],
        "conclusion": [
            "The training modules covered in this report provided meaningful exposure to the supervisory, interpretive, and communication-oriented side of the hostel management system. Through reporting, approvals, analytics, and notifications, the training period strengthened understanding of how data becomes useful for decision making.",
            "The overall learning experience also reinforced the importance of trustworthy summaries, controlled review flow, and timely communication in administrative software systems. These insights make the training highly relevant for future work in systems that support management and oversight.",
        ],
        "appendix_items": [
            "Screenshot of caretaker report page or report action panel exercise",
            "Screenshot of admin approval page or approval card training activity",
            "Screenshot of report generation or review notes",
            "Screenshot of PDF/report support study activity",
            "Screenshot of analytics dashboard or summary screen",
            "Screenshot of count-based administrative observation exercise",
            "Screenshot of notification module or notification record training activity",
            "Screenshot of event-awareness workflow notes",
            "Screenshot of source-data-to-report mapping exercise",
            "Screenshot of approval decision flow diagram or notes",
            "Screenshot of analytics interpretation checklist",
            "Screenshot of backend controller study notes for report modules",
            "Screenshot of notification handling logic review",
            "Screenshot of cross-module dependency notes for analytics",
            "Screenshot of training verification activity for reports",
            "Screenshot of review or summary testing exercise",
            "Screenshot of module integration notes for supervisory workflows",
            "Screenshot of final training checklist for reporting and analytics modules",
        ],
        "appendix_notes": [
            "Screenshots in this appendix should reflect only training work related to reports, approvals, analytics, and notifications.",
            "Avoid using broad project screenshots unless they directly demonstrate the assigned learning modules.",
            "Each figure may briefly mention the reviewed workflow, summary, or communication behavior studied during training.",
        ],
        "support_items": [
            {"material": "Report page screenshot", "purpose": "Shows summary and review training activity"},
            {"material": "Approval screen screenshot", "purpose": "Documents oversight workflow learning"},
            {"material": "Analytics dashboard screenshot", "purpose": "Represents interpretation-focused training"},
            {"material": "Notification workflow note", "purpose": "Shows communication support understanding"},
            {"material": "Dependency mapping notes", "purpose": "Documents system-level analysis"},
        ],
    },
    {
        "member_name": "Pragnya",
        "filename": "Training_Report_Pragnya.docx",
        "report_title": "Authentication, Access Control, Student Signup and Resident Management in MERN Stack",
        "introduction": [
            "This individual training report is prepared for Pragnya and focuses on the training modules related to authentication, access control, student signup flow, and resident management inside the hostel management system. These modules are foundational because they determine who can enter the system, what they can access, and how resident information is created and maintained.",
            "From a training perspective, these modules are especially important because they combine security awareness, workflow review, validation discipline, and role-based system structure. Unlike purely operational record modules, authentication and resident management influence the behavior of almost every other module in the platform.",
            "The report concentrates on learning activities such as understanding login flow, role-based screen routing, signup submission, request review, user and student record handling, and broader resident information management. This created strong exposure to how identity and access form the base layer of a management application.",
            "The following chapters therefore document how these modules were studied, which exercises and observations were used during training, and what technical as well as professional learning outcomes emerged from this domain area.",
        ],
        "objective_intro": "The training objectives for this report were centered on identity flow, controlled access, and resident-related record management. The major objectives were:",
        "objectives": [
            "To understand login, forgot-password, and authentication-related workflow at both frontend and backend levels.",
            "To study role-based access control and how different users are routed to different actions or screens.",
            "To gain practical exposure to student signup requests, review flow, and user creation-related operations.",
            "To understand how resident records are managed, viewed, and maintained in student and admin modules.",
            "To improve knowledge of validation, user identity handling, and controlled access patterns in a hostel system.",
            "To strengthen awareness of how foundational identity modules support the safety and structure of the entire platform.",
        ],
        "domain_overview": [
            "Authentication, access control, signup, and resident management form the foundational trust layer of the system. Without these modules, other operational and financial workflows cannot be applied safely because the platform would lack role clarity, identity verification, and structured resident records.",
            "These modules are valuable from a training perspective because they combine user experience and security concerns. A good authentication module must be simple for genuine users while still preserving access discipline, and resident-management modules must remain organized and reliable because they influence many later workflows.",
        ],
        "tech_blocks": [
            {"title": "MongoDB and Identity Records", "paragraphs": ["User, student, and signup-request models are central to these modules. Through training, it became clear that identity-related records need careful structure because they affect onboarding, role mapping, and resident data reliability."]},
            {"title": "Express.js and Access Logic", "paragraphs": ["Authentication and signup-related controllers illustrate how backend logic controls entry, validation, role handling, and request management. This provided strong training value in understanding secure and structured workflow implementation."]},
            {"title": "React.js for Role-Based Access Screens", "paragraphs": ["Frontend pages such as login, signup, student request review, and admin user screens demonstrate how role-based behavior is made visible to users. These screens helped in understanding how access control becomes practical at UI level."]},
        ],
        "modules": [
            {
                "name": "Authentication and Access Control",
                "frontend": ["LoginPage", "LoginForm", "ForgotPasswordForm"],
                "backend": ["authController.js", "authRoutes.js", "User.js", "PasswordResetOtp.js"],
                "focus": "Login flow, forgot-password support, role access, and identity validation",
                "overview": [
                    "Authentication and Access Control formed one of the most foundational training modules because it determines how the system recognizes users and protects role-specific functionality. Through this module, the training provided meaningful exposure to how software platforms manage entry, validation, and permission-driven navigation.",
                    "A key learning point was that authentication is not only about successful login. It also includes controlled password recovery, protection of restricted actions, and consistent role awareness across screens and backend endpoints. This broadened the understanding of what identity management means in a real application.",
                    "The module also helped illustrate how usability and control must work together. Users should be able to access the right screens smoothly, but the system must still enforce boundaries between student, caretaker, and admin roles. Observing this balance added strong value to the training experience.",
                ],
                "tasks": [
                    "Studied login forms and request handling related to user authentication.",
                    "Reviewed forgot-password workflow and the role of reset-support records.",
                    "Observed how role-based access influences what screens and actions are available to different users.",
                    "Analyzed how backend routes and controllers contribute to secure workflow behavior.",
                ],
                "outcomes": [
                    "This module improved understanding of identity validation, role-based control, and secure workflow design.",
                    "It also strengthened appreciation for the fact that access control is the foundation on which all other system modules depend.",
                ],
                "files": [
                    {"layer": "Frontend", "files": "LoginPage.jsx, LoginForm.jsx, ForgotPasswordForm.jsx", "relevance": "Used to understand user entry and recovery workflow"},
                    {"layer": "Backend", "files": "authController.js, authRoutes.js", "relevance": "Used to study login logic and access-related route behavior"},
                    {"layer": "Models", "files": "User.js, PasswordResetOtp.js", "relevance": "Used to understand identity and recovery-related records"},
                ],
            },
            {
                "name": "Student Signup and Resident Management",
                "frontend": ["StudentSignupForm", "AdminStudentSignupPage", "CaretakerStudentSignupPage", "StudentSignupRequestsPage", "AdminStudentsPage", "AdminUsersPage"],
                "backend": ["studentSignupController.js", "studentController.js", "adminUsersController.js", "authController.js", "StudentSignupRequest.js", "Student.js"],
                "focus": "Signup request capture, review flow, student record handling, and resident/user administration",
                "overview": [
                    "Student Signup and Resident Management formed the second major training area in this report. This module group is important because it handles the creation and maintenance of the resident data that later powers hostel allocation, billing, consumption, and reporting-related workflows.",
                    "Training in this area demonstrated the difference between initial request flow and managed resident records. A signup request is not the same as an approved student record, and user administration is not the same as student profile handling. Understanding these distinctions improved confidence in studying workflow states and administrative responsibilities.",
                    "This module also provided useful exposure to data reliability and record lifecycle management. Resident-related information must remain organized because it influences multiple parts of the system, from access permissions to hostel operations and communication flow.",
                ],
                "tasks": [
                    "Studied student signup form structure and request submission behavior.",
                    "Observed how admin and caretaker screens can review signup requests.",
                    "Reviewed student and user management pages to understand resident record maintenance.",
                    "Analyzed the role of student records in supporting broader hostel workflows.",
                ],
                "outcomes": [
                    "This module improved understanding of onboarding workflow, review-based acceptance, and structured resident record handling.",
                    "It also strengthened awareness of how resident data quality affects many other system modules downstream.",
                ],
                "files": [
                    {"layer": "Frontend", "files": "StudentSignupForm.jsx, AdminStudentSignupPage.jsx, CaretakerStudentSignupPage.jsx, StudentSignupRequestsPage.jsx, AdminStudentsPage.jsx, AdminUsersPage.jsx", "relevance": "Used to understand signup review and resident-management UI"},
                    {"layer": "Backend", "files": "studentSignupController.js, studentController.js, adminUsersController.js, authController.js", "relevance": "Used to study request flow, resident management, and user administration"},
                    {"layer": "Models", "files": "StudentSignupRequest.js, Student.js, User.js", "relevance": "Used to understand onboarding and resident persistence structure"},
                ],
            },
        ],
        "integration": [
            "Authentication, access control, signup, and resident management are tightly connected because each stage builds on the previous one. Users must be identified, roles must be recognized, signup requests must be reviewed appropriately, and resident records must remain structured for later operational use.",
            "These modules also show that foundational workflows can influence the entire system. If access control or resident data is weak, every downstream module becomes less reliable. This made the training especially important for understanding the trust layer of application design.",
        ],
        "integration_map": [
            {"step": "System entry", "primary": "Authentication", "connected": "Access Control", "reason": "Users need validated identity before role-based access is applied"},
            {"step": "Onboarding request", "primary": "Student Signup", "connected": "Admin/Caretaker Review", "reason": "Signup records must be reviewed before becoming trusted resident data"},
            {"step": "Resident creation", "primary": "Resident Management", "connected": "User and student records", "reason": "Approved onboarding must become structured operational identity"},
            {"step": "Role-based operation", "primary": "Access Control", "connected": "All system modules", "reason": "Permissions determine how users interact with operational modules"},
        ],
        "weeks": [
            {"week": "1", "focus": "Login and access orientation", "task": "Studied login flow and role-based route understanding", "outcome": "Built foundation in system entry logic"},
            {"week": "2", "focus": "Recovery workflow awareness", "task": "Reviewed forgot-password support and identity-related records", "outcome": "Improved understanding of secure recovery behavior"},
            {"week": "3", "focus": "Signup request study", "task": "Observed student signup form and request capture flow", "outcome": "Learned onboarding sequence"},
            {"week": "4", "focus": "Review and approval perspective", "task": "Studied admin and caretaker signup review screens", "outcome": "Understood supervised onboarding"},
            {"week": "5", "focus": "Resident record handling", "task": "Reviewed student and user management pages", "outcome": "Improved resident-management awareness"},
            {"week": "6", "focus": "Cross-module trust layer", "task": "Mapped how access and resident records support other modules", "outcome": "Built system-level identity understanding"},
            {"week": "7", "focus": "Validation and discipline", "task": "Analyzed role checks and record correctness needs", "outcome": "Strengthened trust-oriented thinking"},
            {"week": "8", "focus": "Summary and reflection", "task": "Reviewed future learning areas in security and management", "outcome": "Improved confidence in foundational modules"},
        ],
        "week_reflections": [
            "The week-wise progression helped show that identity-related workflows are easier to understand when separated into entry, onboarding, approval, and resident-management stages. This structure improved conceptual clarity during the training period.",
            "A major reflection from this training is that foundational modules often appear simple on the surface, but they carry a large responsibility because they influence permissions, records, and trust throughout the rest of the application.",
        ],
        "skills": [
            {"title": "Technical Skills", "bullets": ["Improved understanding of authentication, role handling, and access-aware workflow", "Better awareness of student signup review and resident record lifecycle", "Increased confidence in studying foundational user and identity modules"], "paragraphs": ["This report area strengthened understanding of the base-layer modules that make the rest of the platform usable and safe."]},
            {"title": "Tools and Software", "bullets": ["Login and signup frontend screens", "Backend controllers for identity and resident management", "MongoDB models for user, student, and request records"], "paragraphs": []},
            {"title": "Soft Skills", "bullets": ["Greater care while studying sensitive modules", "Improved structured thinking around permissions and role clarity", "Better communication of foundational workflow dependencies"], "paragraphs": []},
        ],
        "industry_practices": [
            "The training in these modules highlighted an important industry principle: identity and access are not optional utilities but foundational requirements. Reliable onboarding, clear role mapping, and disciplined resident record management support the entire software environment.",
            "Another useful insight was that security-related modules must remain understandable to users while still maintaining control. This balance between usability and protection is a key part of professional application development.",
        ],
        "challenges": [
            "One challenge was understanding how many different responsibilities exist within the identity layer, including login, access, recovery, signup, approval, and resident management.",
            "Another challenge was recognizing the distinction between request-level onboarding data and trusted resident records used in regular operations.",
            "These challenges were addressed by studying the modules in lifecycle order and then mapping how they influence broader system trust and role behavior.",
        ],
        "key_learnings": [
            "Authentication and access control shape the trust layer of the entire platform.",
            "Signup requests and resident records should be treated as different workflow stages.",
            "Role-based design improves clarity and safety in multi-user systems.",
            "Resident data quality influences many downstream hostel operations.",
        ],
        "self_assessment": [
            "The training period improved confidence in understanding how identity, permissions, and resident records are structured in a practical system. This was especially useful because these modules influence nearly every other workflow area in the platform.",
            "Future learning is still needed in advanced authentication strategies, deeper security patterns, and larger-scale user administration. However, the current training has created strong practical clarity around the foundational modules covered in this report.",
        ],
        "assessment_table": [
            {"area": "Authentication understanding", "level": "Good working clarity", "future": "Study deeper security patterns"},
            {"area": "Role-based access awareness", "level": "Comfortable with practical flow", "future": "Explore more advanced permission systems"},
            {"area": "Signup workflow knowledge", "level": "Good conceptual understanding", "future": "Study richer onboarding validation"},
            {"area": "Resident management", "level": "Working understanding", "future": "Improve administrative record lifecycle depth"},
        ],
        "conclusion": [
            "The training modules documented in this report provided meaningful exposure to the foundational trust and resident-management layer of the hostel management system. Through authentication, access control, signup, and resident workflows, the training period strengthened understanding of how system structure begins with identity and role clarity.",
            "The experience also reinforced the importance of disciplined onboarding, controlled access, and reliable resident records in software systems that support many user roles. These insights make the training highly relevant for future work in secure and structured web applications.",
        ],
        "appendix_items": [
            "Screenshot of login form or authentication training activity",
            "Screenshot of forgot-password or recovery support exercise",
            "Screenshot of role-based routing or access observation",
            "Screenshot of student signup form training activity",
            "Screenshot of signup request review page",
            "Screenshot of caretaker or admin review workflow for requests",
            "Screenshot of student management page or resident record view",
            "Screenshot of user management or role-handling screen",
            "Screenshot of onboarding lifecycle notes from request to resident record",
            "Screenshot of backend controller study notes for auth modules",
            "Screenshot of user/student model review activity",
            "Screenshot of validation or permission checklist",
            "Screenshot of route protection or role-awareness notes",
            "Screenshot of training verification activity for signup flow",
            "Screenshot of resident-management module relationship notes",
            "Screenshot of trust-layer mapping across system roles",
            "Screenshot of final checklist for authentication and resident modules",
            "Screenshot of module integration notes for onboarding and access control",
        ],
        "appendix_notes": [
            "Include only screenshots that directly represent training work related to authentication, access control, signup, or resident management.",
            "Avoid unrelated project screenshots unless they clearly support the assigned identity or resident-management module context.",
            "Each figure may briefly explain the security, onboarding, or role-handling aspect studied during training.",
        ],
        "support_items": [
            {"material": "Login/auth screenshot", "purpose": "Shows entry and access-related training activity"},
            {"material": "Signup request screenshot", "purpose": "Represents onboarding workflow understanding"},
            {"material": "Resident management screenshot", "purpose": "Documents resident record study"},
            {"material": "Controller/model notes", "purpose": "Shows backend identity-layer analysis"},
            {"material": "Permission checklist", "purpose": "Captures access-control learning summary"},
        ],
    },
]


EXTRA_APPENDIX_SUFFIXES = [
    "workflow sequence mapping exercise",
    "validation checklist used during training",
    "module integration observation notes",
    "backend route or controller review activity",
    "data model or schema training notes",
    "UI refinement or form-structure practice",
    "testing or verification exercise for the assigned module",
    "final revision checklist prepared during training",
    "sample data review exercise",
    "role-based interaction observation",
    "error-handling or debugging note",
    "module summary prepared for weekly review",
    "input-output flow mapping activity",
    "record status tracking exercise",
    "documentation note prepared during training",
    "final understanding recap diagram",
]


for member in MEMBERS:
    primary_label = member["modules"][0]["name"]
    extra_items = [f"Screenshot of {primary_label.lower()} {suffix}" for suffix in EXTRA_APPENDIX_SUFFIXES]
    member["appendix_items"].extend(extra_items)


def main():
    for member in MEMBERS:
        out_path = build_report(member)
        print(out_path)


if __name__ == "__main__":
    main()

from datetime import datetime
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


TITLE = "FinTrix: Web-Based Hostel Management System"
SUBTITLE = "Synopsis Report"
OUTPUT_PATH = "docs/FinTrix_Synopsis_Report.docx"


REFERENCES = [
    "[1] J. Oyeniyi, \"Development of Web-Based Hostel Management System,\" British Journal of Computer, Networking and Information Technology, vol. 8, no. 1, pp. 30-41, 2025.",
    "[2] D. Narkhede, R. Bamgude, M. Sonawane, and M. Shevade, \"Hostel Management System (HMS),\" International Journal for Research in Applied Science & Engineering Technology, vol. 10, issue IV, 2022.",
    "[3] A. Pawar, V. Ukarande, S. Kale, P. Kshirsagar, S. G. Ekdante, and J. M. Shaikh, \"Hostel Management System,\" International Journal of Recent Research in Mathematics Computer Science and Information Technology, vol. 12, issue 1, pp. 47-56, 2025.",
    "[4] Dinesh B., Gogul Nithin R., Pavatharani R., Sneha R., and C. Senthilkumar, \"Implementation of Hostel Management with Automation Using Design Thinking,\" IJCRT, vol. 10, issue 4, 2022.",
    "[5] A. M. Diyaolu, O. B. Abodunrin, A. A. Adedamola, R. S. Ogunode, and O. Omoloba, \"Development of an E-Based Hostel Management System,\" International Journal of Innovative Science and Research Technology, vol. 9, issue 6, 2024.",
    "[6] Aravinth M., Nithin K., and J. Kayalvizhi, \"Automated Hostel Management System,\" International Journal of Scientific Research & Engineering Trends, vol. 11, issue 1, 2025.",
    "[7] R. K. Bista, A. J. Karki, B. V. M. Reddy, U. Aakash, R. A. Makaram, and S. Das, \"Hostel Management System,\" International Journal of Trend in Scientific Research and Development, vol. 2, issue 4, 2018.",
    "[8] T. Sai Prasad Reddy, S. Jayakrishna, I. Vasu, N. Leela Siddhiswar, and G. Raviteja, \"Hostel Management System,\" Journal of Emerging Technologies and Innovative Research, vol. 13, issue 4, 2026.",
    "[9] Likhin S. and Shashidhar Kini K., \"Hostel Management System,\" Journal of Emerging Technologies and Innovative Research, vol. 12, issue 7, 2025.",
]


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.color.rgb = RGBColor(31, 78, 121)
    if level == 1:
        run.font.size = Pt(15)
    elif level == 2:
        run.font.size = Pt(12.5)
    return p


def add_body_paragraph(doc, text, first=False):
    p = doc.add_paragraph()
    p.style = "Normal"
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    if first:
        p.paragraph_format.first_line_indent = Inches(0.3)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)


def add_table_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(68, 68, 68)


def add_label_paragraph(doc, label, value):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(f"{label}: ")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    value_run = p.add_run(value)
    value_run.font.name = "Calibri"
    value_run.font.size = Pt(11)


def build_cover_page(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Inches(1.0)
    run = p.add_run(SUBTITLE)
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(31, 78, 121)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run(TITLE)
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(46, 46, 46)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run("Prepared as a project synopsis for a modern hostel administration and student service platform")
    run.italic = True
    run.font.name = "Calibri"
    run.font.size = Pt(12)

    meta_lines = [
        "Project Type: Academic Synopsis / System Proposal",
        "Application Domain: Hostel Administration, Billing, and Student Services",
        "Implementation Context: Full-stack web application",
        "Reference Basis: Attached synopsis template and hostel management research papers",
        f"Prepared On: {datetime.now().strftime('%d %B %Y')}",
    ]

    for line in meta_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(line)
        run.font.name = "Calibri"
        run.font.size = Pt(11)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(18)
    run = note.add_run(
        "This report proposes a role-based web system that digitizes student registration, hostel allocation, billing, payments, reporting, and operational monitoring."
    )
    run.font.name = "Calibri"
    run.font.size = Pt(11)


def add_section_break(doc):
    doc.add_section(WD_SECTION.NEW_PAGE)


def build_document():
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)

    build_cover_page(doc)
    add_section_break(doc)

    add_heading(doc, "Abstract", 1)
    add_body_paragraph(
        doc,
        "FinTrix is a web-based hostel management system designed to replace fragmented, manual, and paper-driven hostel administration with a secure digital workflow. The proposed system centralizes student onboarding, hostel and room data, mess and utility consumption, charge management, bill generation, payment tracking, approvals, analytics, notifications, and report generation inside a single platform. Existing studies consistently show that hostel administration suffers from delayed records, weak transparency, inefficient room allocation, poor financial traceability, and limited access to decision-ready information [1], [2], [5]. In response, this synopsis proposes a full-stack role-based application where students, caretakers, wardens, and administrators interact through dedicated dashboards. The architecture combines a React-based user interface, an Express.js service layer, MongoDB data persistence, JWT-based authentication, validation middleware, reporting utilities, and scheduled billing jobs. The expected outcome is a scalable system that improves operational efficiency, reduces clerical error, strengthens accountability, and provides a better service experience to both hostel staff and students.",
        first=True,
    )

    add_heading(doc, "1. Introduction", 1)
    add_body_paragraph(
        doc,
        "Hostels attached to colleges and universities manage a large volume of operational activities every day, including admissions, room allotment, occupancy tracking, fee collection, mess expenses, utility billing, complaints, and communication with residents. When these functions are handled through registers, spreadsheets, and isolated files, data becomes inconsistent, retrieval becomes slow, and administrative follow-up becomes difficult. Researchers have repeatedly argued that web-based hostel systems improve service coordination by making critical information available in real time and by reducing dependence on manual bookkeeping [1], [3], [6].",
        first=True,
    )
    add_body_paragraph(
        doc,
        "The proposed project, FinTrix, extends the idea of a traditional hostel management system by placing strong emphasis on expense tracking, mess bill computation, advance adjustments, guest charges, payment monitoring, analytics, and downloadable reports. This makes the platform suitable not only for resident record management but also for operational finance and accountability. Literature on hostel automation emphasizes that digitization should not merely store data; it should support transparency, traceability, and faster decision-making for administrators [4], [5], [9].",
        first=True,
    )
    add_body_paragraph(
        doc,
        "FinTrix is conceived as a multi-role web application in which each stakeholder sees only the features relevant to their responsibilities. Students can review bills and payment status, caretakers can manage hostel operations and recurring charges, and administrators can supervise hostels, users, approvals, and analytics. This role-based structure aligns with modern hostel administration practices and supports better governance through controlled access and auditable workflows.",
        first=True,
    )

    add_heading(doc, "2. Problem Statement", 1)
    add_body_paragraph(
        doc,
        "Many hostel environments still rely on notebooks, disconnected spreadsheets, informal approval chains, and repetitive calculations for day-to-day administration. This creates several operational problems: duplicate records, delayed updates, difficulty in tracing dues, weak coordination between hostel staff and administrators, and limited visibility into student-level expenses. Prior work highlights that manual hostel processes often become inaccurate and time-consuming as the number of residents grows [2], [3], [7].",
        first=True,
    )
    add_body_paragraph(
        doc,
        "A more specific challenge appears in hostels that manage shared mess expenses, electricity or other utility bills, student advances, and multiple incidental charges. In such cases, administrators must perform repeated calculations, maintain evidence for every adjustment, and answer student queries about fairness and correctness. Without a centralized system, the chances of billing disputes, omitted records, and reporting delays increase significantly. Therefore, the core problem is the absence of a unified, secure, and transparent web platform that integrates hostel administration with resident financial management and reporting.",
        first=True,
    )

    add_heading(doc, "3. Scope", 1)
    add_body_paragraph(
        doc,
        "The scope of the proposed system covers the digital management of hostel-related academic and operational workflows within an institutional environment. The system will support student registration and approval, hostel data management, resident records, room and hostel allocation support, mess and consumption entry, electricity or utility tracking, fixed and variable charge handling, advance management, payment history, notifications, reports, and analytics. The application is intended for use through a web browser and can serve students, caretakers, wardens, and administrators.",
        first=True,
    )
    add_body_paragraph(
        doc,
        "The synopsis does not target IoT-enabled room monitoring, biometric gate hardware integration, or fully automated seat allocation algorithms at this stage. Instead, it focuses on building a reliable, modular, and scalable software foundation that can later integrate such features. This boundary keeps the project feasible while still delivering a meaningful improvement over manual and semi-digital hostel administration.",
        first=True,
    )

    add_heading(doc, "4. Objectives", 1)
    add_bullet(doc, "To design and implement a centralized web-based hostel management platform for operational and financial workflows.")
    add_bullet(doc, "To reduce manual paperwork, repetitive calculations, and record inconsistencies across hostel activities.")
    add_bullet(doc, "To provide role-based access for students, caretakers, wardens, and administrators.")
    add_bullet(doc, "To automate bill preparation, charge tracking, payment monitoring, and report generation.")
    add_bullet(doc, "To improve transparency, accountability, and auditability through validated and traceable data handling.")
    add_bullet(doc, "To generate actionable insights through dashboards and analytics for faster administrative decisions.")

    add_heading(doc, "5. Literature Survey", 1)
    add_body_paragraph(
        doc,
        "Oyeniyi's 2025 work presents a web-based hostel management system that emphasizes registration, room allocation, maintenance, reporting, and the value of real-time administrative access [1]. The paper is especially relevant because it frames hostel management as an information coordination problem, not just a storage problem. That perspective supports the FinTrix design choice of combining operational modules with dashboards, approvals, and reports.",
        first=True,
    )
    add_body_paragraph(
        doc,
        "Narkhede et al. describe an online hostel management system intended to reduce the strain created by growing institutional hostel populations and manual workflows [2]. A similar concern is visible in Pawar et al., where the focus is on replacing manual hostel administration with a more website-oriented system to improve efficiency [3]. These studies collectively justify the need for digitization and reinforce the proposed shift from paper-based processes to a browser-accessible platform.",
        first=True,
    )
    add_body_paragraph(
        doc,
        "Dinesh et al. approach hostel automation using design thinking and highlight the need for systems that reflect user pain points while making operations more transparent and manageable [4]. Diyaolu et al. extend the discussion by presenting an e-based solution that improves access to administrative information and enhances service delivery [5]. Together, these works suggest that a successful hostel system must balance usability, process automation, and visibility for decision makers.",
        first=True,
    )
    add_body_paragraph(
        doc,
        "Recent publications from 2025 and 2026 continue to emphasize automation, financial traceability, and operational streamlining. Aravinth and colleagues discuss an automated hostel management system that covers student registration, room allocation, and mess management [6]. The newer JETIR studies also underline the continuing importance of centralized records and faster hostel administration in educational institutions [8], [9]. However, many surveyed systems remain broad at a conceptual level and give less attention to recurring financial controls such as advances, detailed consumption-based billing, guest charges, downloadable report flows, and role-specific dashboards. FinTrix aims to address that gap by combining resident administration with expense-aware hostel finance management.",
        first=True,
    )

    add_heading(doc, "6. Requirement Specification", 1)
    add_heading(doc, "6.1 Functional Requirements", 2)
    functional_items = [
        "User authentication and authorization for student, caretaker, dean, warden, and admin roles.",
        "Student signup request submission, review, approval, and rejection workflows.",
        "Hostel and resident record management with update and search support.",
        "Expense, hostel expense, utility consumption, charge, and guest charge entry modules.",
        "Bill generation, bill breakdown access, and payment tracking facilities.",
        "Advance management and adjustment visibility for student accounts.",
        "Analytics dashboards and report generation with downloadable outputs.",
        "Notification support for status updates and administrative communication.",
    ]
    for item in functional_items:
        add_bullet(doc, item)

    add_heading(doc, "6.2 Non-Functional Requirements", 2)
    non_functional_items = [
        "Security through token-based authentication, password hashing, request sanitization, and protected routes.",
        "Usability through role-specific navigation, dashboard-based interaction, and responsive web access.",
        "Scalability through modular routes, controllers, and database models.",
        "Maintainability through separation of frontend, backend, middleware, and model layers.",
        "Reliability through validation, centralized error handling, and scheduled recurring bill jobs.",
        "Performance through API-based data access and optimized retrieval of operational records.",
    ]
    for item in non_functional_items:
        add_bullet(doc, item)

    add_heading(doc, "7. System Requirements", 1)
    add_table_title(doc, "Recommended deployment and usage requirements")
    req_rows = [
        ("Client Device", "Dual-core CPU with 4 GB RAM minimum; modern laptop or desktop with 8 GB RAM recommended."),
        ("Browser", "Google Chrome, Microsoft Edge, or Mozilla Firefox; the latest stable version is preferred."),
        ("Server", "At least 4 CPU cores and 8 GB RAM for small deployments; 8 CPU cores and 16 GB RAM for institutional use."),
        ("Database", "MongoDB with regular backup support; a managed or cloud-backed cluster is recommended for reliability."),
        ("Network", "Stable LAN or broadband access for staff and students."),
        ("Operating System", "Windows, Linux, or macOS for development; Linux is preferred for production deployment."),
    ]
    for label, value in req_rows:
        add_label_paragraph(doc, label, value)

    add_heading(doc, "8. Tools and Technology", 1)
    add_table_title(doc, "Proposed technology stack")
    tech_rows = [
        ("Frontend", "React with Vite for a responsive single-page user interface and faster development feedback."),
        ("State and Data", "Zustand, TanStack React Query, and Axios for session state, API fetching, caching, and HTTP communication."),
        ("Backend", "Node.js with Express.js to implement REST endpoints, request flow control, and business logic."),
        ("Database", "MongoDB with Mongoose for storing hostel, student, expense, bill, and payment records."),
        ("Security", "JWT, bcrypt, helmet, and express-rate-limit for authentication, password protection, secure headers, and abuse control."),
        ("Documentation and Output", "Swagger and PDFKit for API documentation, report generation, and downloadable bill outputs."),
        ("Automation", "node-cron for scheduled billing operations."),
        ("Testing", "Vitest, Jest, and Supertest for frontend and backend quality validation."),
    ]
    for label, value in tech_rows:
        add_label_paragraph(doc, label, value)

    add_heading(doc, "9. System Architecture", 1)
    add_body_paragraph(
        doc,
        "FinTrix follows a layered web architecture. The presentation layer is a React-based frontend that exposes dedicated dashboards and workflow pages for students, caretakers, and administrators. The application layer is built with Express.js and organizes logic into routes, controllers, middleware, scheduled jobs, and utility services. The data layer uses MongoDB through Mongoose models to store students, hostels, expenses, charges, payments, reports, notifications, and related entities. Security and validation are handled through authentication middleware, request sanitization, and input validation before data is committed.",
        first=True,
    )
    add_table_title(doc, "Architectural view of the proposed system")
    arch_rows = [
        ("Presentation Layer", "React pages, protected routes, dashboards, and forms collect user input and present role-specific workflows."),
        ("API Layer", "Express routes handle authentication, hostel records, students, billing, expenses, analytics, reports, and notifications."),
        ("Business Logic Layer", "Controllers, middleware, validators, cron jobs, and utility modules implement rules, calculations, and approvals."),
        ("Data Layer", "Mongoose models persist users, students, hostels, charges, expenses, payments, reports, and notifications in MongoDB."),
        ("Security Layer", "JWT authentication, hashing, sanitization, secure headers, and rate limiting protect the application."),
    ]
    for label, value in arch_rows:
        add_label_paragraph(doc, label, value)

    add_body_paragraph(
        doc,
        "The role flow can be summarized as follows: students log in to review bills, payments, and notifications; caretakers enter operational and financial records; wardens and deans oversee approvals and analytics; and administrators manage global hostel configuration, users, and institutional monitoring. This structure improves accountability because every major activity passes through authenticated and role-constrained interfaces.",
        first=True,
    )

    add_heading(doc, "10. Expected Outcomes", 1)
    outcomes = [
        "A functioning web-based hostel management platform with authenticated multi-role access.",
        "Faster and more accurate management of registration, billing, expenses, and payments.",
        "Improved transparency in hostel finance through bill breakdowns, advances, and reports.",
        "Reduced paperwork and lower dependency on manual calculation and record lookup.",
        "Better administrative decisions through summary dashboards and analytics.",
        "A scalable software base for future additions such as complaint management, room allocation automation, and mobile access.",
    ]
    for item in outcomes:
        add_bullet(doc, item)

    add_heading(doc, "11. Conclusion", 1)
    add_body_paragraph(
        doc,
        "The proposed FinTrix system addresses a practical and recurring institutional problem: the difficulty of managing hostel operations and hostel-related finances through manual and disconnected methods. Existing research strongly supports the transition toward centralized and web-based hostel platforms because they improve efficiency, transparency, and service quality [1], [4], [5]. Building on those findings, this synopsis presents a realistic full-stack architecture that combines resident management with expense-aware billing, approval workflows, analytics, and secure access control.",
        first=True,
    )
    add_body_paragraph(
        doc,
        "By aligning the system design with real operational needs such as charges, consumption, advances, payments, and role-based oversight, the proposed project goes beyond a basic hostel record system. It aims to become a dependable administrative tool that can support both day-to-day execution and longer-term institutional accountability. Therefore, the development of FinTrix is technically feasible, academically relevant, and operationally valuable for modern hostel environments.",
        first=True,
    )

    add_heading(doc, "12. References", 1)
    for ref in REFERENCES:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(ref)
        run.font.name = "Calibri"
        run.font.size = Pt(10.5)

    doc.save(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
